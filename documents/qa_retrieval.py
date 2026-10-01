"""Truy xuất đoạn văn liên quan câu hỏi — toàn văn tài liệu + hướng dẫn sử dụng.

Văn bản trích từ file được cache theo (id, updated_at) nên chỉ đọc file một lần
cho mỗi phiên bản tài liệu.
"""

from __future__ import annotations

import hashlib
import io
import logging
import math
import os
import re
import unicodedata
from collections import Counter
from dataclasses import dataclass, field

from django.core.cache import cache
from django.utils.html import strip_tags

logger = logging.getLogger(__name__)

DOC_TEXT_CACHE_TIMEOUT = 60 * 60 * 24 * 30
GUIDE_CACHE_TIMEOUT = 60 * 60 * 6
MAX_FILE_BYTES = 40 * 1024 * 1024
MAX_PDF_PAGES = 300
MAX_SHEET_ROWS = 2000
MAX_DOC_CHARS = 300_000

CHUNK_CHARS = 1100
CHUNK_OVERLAP = 180
MAX_CHUNKS_PER_SOURCE = 4

BM25_K1 = 1.4
BM25_B = 0.72

STOPWORDS = {
    'toi', 'ban', 'la', 'gi', 'co', 'khong', 'duoc', 'the', 'nao', 'va', 'cua',
    'trong', 'tren', 'portal', 'justplay', 'xin', 'cho', 'hay', 've', 'mot', 'cac',
    'giup', 'gui', 'link', 'nhu', 'thi', 'ma', 'de', 'o', 'voi', 'nay', 'do', 'khi',
    'neu', 'can', 'muon', 'lam', 'sao', 'a', 'oi', 'nhe', 'vay', 'roi', 'da', 'se',
    'dang', 'nhung', 'nhieu', 'it', 'minh', 'em', 'anh', 'chi', 'u', 'ah', 'vui', 'long',
}

_HTML_BLOCK_RE = re.compile(r'</?(p|div|li|ul|ol|h[1-6]|tr|table|br|section|article)[^>]*>', re.I)
_WORD_RE = re.compile(r'\w+', re.UNICODE)


def strip_accents(text: str) -> str:
    normalized = unicodedata.normalize('NFD', text or '')
    out = ''.join(ch for ch in normalized if unicodedata.category(ch) != 'Mn')
    return out.replace('đ', 'd').replace('Đ', 'D')


def tokenize(text: str) -> list[str]:
    words = [w for w in _WORD_RE.findall(strip_accents((text or '').lower())) if w not in STOPWORDS]
    words = [w for w in words if len(w) > 1 or w.isdigit()]
    bigrams = [f'{a}_{b}' for a, b in zip(words, words[1:])]
    return words + bigrams


def html_to_text(html: str) -> str:
    text = _HTML_BLOCK_RE.sub('\n', html or '')
    text = strip_tags(text)
    text = re.sub(r'&nbsp;', ' ', text)
    text = re.sub(r'[ \t]+', ' ', text)
    text = re.sub(r'\n\s*\n+', '\n\n', text)
    text = re.sub(r'(?m)^\s*(\d{1,2})\s*\n+(?=\S)', r'\1. ', text)
    return text.strip()


# --- Trích văn bản từ file ---


def _read_file_bytes(fieldfile) -> bytes | None:
    try:
        size = fieldfile.size
    except Exception:
        size = None
    if size and size > MAX_FILE_BYTES:
        logger.info('QA skip large file %s (%s bytes)', fieldfile.name, size)
        return None
    try:
        with fieldfile.open('rb') as fh:
            return fh.read()
    except Exception:
        logger.warning('QA cannot open file %s', getattr(fieldfile, 'name', ''), exc_info=True)
        return None


def _extract_pdf(data: bytes) -> str:
    try:
        import fitz

        with fitz.open(stream=data, filetype='pdf') as pdf:
            return '\n'.join(page.get_text() for page in pdf.pages(0, min(len(pdf), MAX_PDF_PAGES)))
    except Exception:
        logger.debug('PyMuPDF failed, fallback pypdf', exc_info=True)
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    return '\n'.join((page.extract_text() or '') for page in reader.pages[:MAX_PDF_PAGES])


def _extract_docx(data: bytes) -> str:
    import docx

    document = docx.Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(' | '.join(cells))
    return '\n'.join(parts)


def _extract_xlsx(data: bytes) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    parts = []
    try:
        for ws in wb.worksheets:
            parts.append(f'[Sheet: {ws.title}]')
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i >= MAX_SHEET_ROWS:
                    break
                cells = [str(v).strip() for v in row if v not in (None, '')]
                if cells:
                    parts.append(' | '.join(cells))
    finally:
        wb.close()
    return '\n'.join(parts)


def extract_file_text(fieldfile, filename: str = '') -> str:
    if not fieldfile:
        return ''
    ext = os.path.splitext((filename or fieldfile.name or '').lower())[1]
    if ext not in {'.pdf', '.docx', '.xlsx', '.xlsm', '.txt', '.md', '.csv'}:
        return ''
    data = _read_file_bytes(fieldfile)
    if not data:
        return ''
    try:
        if ext == '.pdf':
            return _extract_pdf(data)
        if ext == '.docx':
            return _extract_docx(data)
        if ext in {'.xlsx', '.xlsm'}:
            return _extract_xlsx(data)
        return data.decode('utf-8', errors='ignore')
    except Exception:
        logger.warning('QA text extraction failed: %s', fieldfile.name, exc_info=True)
        return ''


def document_full_text(doc) -> str:
    """Toàn văn tài liệu (nội dung soạn + file đính kèm), cache theo phiên bản."""
    stamp = int(doc.updated_at.timestamp()) if doc.updated_at else 0
    key = f'qa:doctext:v1:{doc.pk}:{stamp}'
    cached = cache.get(key)
    if cached is not None:
        return cached

    parts = []
    if doc.summary:
        parts.append(doc.summary)
    if doc.body:
        parts.append(html_to_text(doc.body))
    seen = set()
    for f, name in ((doc.pdf_file, ''), (doc.original_file, doc.original_filename)):
        if f and f.name not in seen:
            seen.add(f.name)
            parts.append(extract_file_text(f, name))
    text = '\n\n'.join(p for p in parts if p and p.strip())[:MAX_DOC_CHARS]
    cache.set(key, text, DOC_TEXT_CACHE_TIMEOUT)
    return text


# --- Chia đoạn & chấm điểm ---


@dataclass
class Passage:
    source: str
    url: str
    text: str
    tokens: Counter = field(default_factory=Counter)
    length: int = 0


def split_passages(text: str) -> list[str]:
    text = re.sub(r'[ \t]+', ' ', text or '').strip()
    if not text:
        return []
    paragraphs = [p.strip() for p in re.split(r'\n+', text) if p.strip()]
    chunks, buf = [], ''
    for para in paragraphs:
        while len(para) > CHUNK_CHARS:
            head, para = para[:CHUNK_CHARS], para[CHUNK_CHARS - CHUNK_OVERLAP:]
            if buf:
                chunks.append(buf)
                buf = ''
            chunks.append(head)
        if len(buf) + len(para) + 1 > CHUNK_CHARS and buf:
            chunks.append(buf)
            buf = buf[-CHUNK_OVERLAP:] + '\n' + para
        else:
            buf = f'{buf}\n{para}' if buf else para
    if buf:
        chunks.append(buf)
    return chunks


_PASSAGE_CACHE: dict[str, list[Passage]] = {}
_PASSAGE_CACHE_MAX = 400


def build_passages(cache_key: str, source: str, url: str, text: str) -> list[Passage]:
    hit = _PASSAGE_CACHE.get(cache_key)
    if hit is not None:
        return hit
    title_tokens = tokenize(source)
    passages = []
    for chunk in split_passages(text):
        toks = Counter(tokenize(chunk))
        for t in title_tokens:
            toks[t] += 2
        passages.append(Passage(source=source, url=url, text=chunk, tokens=toks, length=sum(toks.values())))
    if len(_PASSAGE_CACHE) >= _PASSAGE_CACHE_MAX:
        _PASSAGE_CACHE.clear()
    _PASSAGE_CACHE[cache_key] = passages
    return passages


def rank_passages(passages: list[Passage], query: str, *, char_budget: int) -> list[Passage]:
    q_tokens = set(tokenize(query))
    if not q_tokens or not passages:
        return []
    n = len(passages)
    avg_len = sum(p.length for p in passages) / n or 1
    df = Counter()
    for p in passages:
        for t in q_tokens:
            if t in p.tokens:
                df[t] += 1

    scored = []
    for p in passages:
        score = 0.0
        for t in q_tokens:
            tf = p.tokens.get(t, 0)
            if not tf:
                continue
            idf = math.log(1 + (n - df[t] + 0.5) / (df[t] + 0.5))
            weight = 1.6 if '_' in t else 1.0
            score += weight * idf * tf * (BM25_K1 + 1) / (tf + BM25_K1 * (1 - BM25_B + BM25_B * p.length / avg_len))
        if score > 0:
            scored.append((score, p))
    scored.sort(key=lambda x: -x[0])
    if not scored:
        return []

    top_score = scored[0][0]
    picked, per_source, used = [], Counter(), 0
    for score, p in scored:
        if score < top_score * 0.18:
            break
        if per_source[p.source] >= MAX_CHUNKS_PER_SOURCE:
            continue
        if used + len(p.text) > char_budget:
            continue
        picked.append(p)
        per_source[p.source] += 1
        used += len(p.text)
    return picked


# --- Nguồn: tài liệu & hướng dẫn ---


def document_passages(docs_with_urls) -> list[Passage]:
    out = []
    for doc, category_name, url in docs_with_urls:
        stamp = int(doc.updated_at.timestamp()) if doc.updated_at else 0
        text = document_full_text(doc)
        out.extend(build_passages(
            f'doc:{doc.pk}:{stamp}',
            f'Tài liệu «{doc.title}» ({category_name})',
            url,
            text,
        ))
    return out


def guide_sections_text(request, guide_url: str) -> list[tuple[str, str, str]]:
    """[(tiêu đề mục, url, văn bản)] — chỉ mục user được xem."""
    from hrm.guide_editor import (
        normalize_section_overrides,
        render_section_inner_default,
        section_display_title,
    )
    from hrm.guide_sections import get_guide_admin_section_ids, get_visible_guide_sections
    from hrm.models import UserGuide

    user = request.user
    guide = UserGuide.load()
    visible = get_visible_guide_sections(user)
    admin_ids = get_guide_admin_section_ids(user)
    stamp = int(guide.updated_at.timestamp()) if guide.updated_at else 0
    sig = hashlib.md5(
        ('|'.join(s['id'] for s in visible) + '#' + '|'.join(sorted(admin_ids))).encode()
    ).hexdigest()
    key = f'qa:guide:v1:{stamp}:{sig}'
    cached = cache.get(key)
    if cached is not None:
        return cached

    from hrm.views_guide import _guide_context

    ctx = _guide_context(request, guide, can_edit=False)
    overrides = normalize_section_overrides(guide.section_overrides)
    out = []
    for sec in ctx['visible_sections']:
        body = (overrides.get(sec['id']) or {}).get('body') or render_section_inner_default(
            sec['id'], request, context=ctx,
        )
        text = html_to_text(body)
        if text:
            title = section_display_title(sec, overrides)
            out.append((title, f'{guide_url}#{sec["id"]}', text))
    cache.set(key, out, GUIDE_CACHE_TIMEOUT)
    return out


def guide_passages(request, guide_url: str) -> list[Passage]:
    out = []
    for title, url, text in guide_sections_text(request, guide_url):
        key = 'guide:' + hashlib.md5(f'{url}\n{text}'.encode()).hexdigest()
        out.extend(build_passages(key, f'Hướng dẫn sử dụng — {title}', url, text))
    return out
