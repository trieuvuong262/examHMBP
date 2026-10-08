"""Import / file mẫu / xuất Excel cho Tiến độ: Tính năng, Mô tả, User flow, Feedback."""

from __future__ import annotations

import html
import io
import re
import unicodedata
import warnings
from dataclasses import dataclass

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

from .models import TienDoItem

MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_ROWS = 1000
HEADER_SCAN_ROWS = 10

TEMPLATE_HEADERS = ('Tính năng', 'Mô tả', 'User flow')

EXPORT_HEADERS = (
    'STT',
    'Tính năng',
    'Mô tả',
    'User flow',
    'Người gửi',
    'Thời gian',
    'Feedback',
    'Ghi chú',
    'Cập nhật',
)

# Tên tiêu đề cột (đã bỏ dấu, viết thường) → field model.
_HEADER_ALIASES = {
    'feature': {'tinh nang', 'ten tinh nang', 'chuc nang', 'feature'},
    'description': {'mo ta', 'description'},
    'user_flow': {'user flow', 'userflow', 'luong', 'luong su dung', 'flow'},
}


class TienDoImportError(Exception):
    pass


@dataclass
class ParsedRow:
    feature: str
    description: str
    user_flow: str


def _norm(value) -> str:
    text = str(value or '').strip().lower().replace('đ', 'd')
    text = unicodedata.normalize('NFD', text)
    text = ''.join(ch for ch in text if unicodedata.category(ch) != 'Mn')
    return ' '.join(text.replace('_', ' ').split())


def _cell_text(value) -> str:
    if value is None:
        return ''
    text = str(value).replace('\r\n', '\n').replace('\r', '\n')
    return text.strip()


def _detect_header(rows):
    """Tìm dòng tiêu đề trong vài dòng đầu → (index dòng, {field: cột})."""
    for row_idx, row in enumerate(rows[:HEADER_SCAN_ROWS]):
        mapping = {}
        for col_idx, value in enumerate(row):
            key = _norm(value)
            for field, aliases in _HEADER_ALIASES.items():
                if key in aliases and field not in mapping:
                    mapping[field] = col_idx
        if 'feature' in mapping:
            return row_idx, mapping
    return None, None


def parse_workbook(file_obj) -> list[ParsedRow]:
    try:
        with warnings.catch_warnings():
            warnings.simplefilter('ignore')
            wb = openpyxl.load_workbook(file_obj, data_only=True, read_only=True)
    except Exception as exc:  # noqa: BLE001 — openpyxl ném nhiều loại lỗi
        raise TienDoImportError('Không đọc được file Excel. Hãy dùng file .xlsx.') from exc

    try:
        ws = wb.active
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
    finally:
        wb.close()

    header_idx, mapping = _detect_header(rows)
    if header_idx is None:
        raise TienDoImportError('Không thấy cột "Tính năng" trong file. Hãy dùng file mẫu.')

    def pick(row, field):
        col = mapping.get(field)
        if col is None or col >= len(row):
            return ''
        return _cell_text(row[col])

    parsed = []
    for row in rows[header_idx + 1:]:
        item = ParsedRow(
            feature=pick(row, 'feature')[:255],
            description=pick(row, 'description'),
            user_flow=pick(row, 'user_flow'),
        )
        if not (item.feature or item.description or item.user_flow):
            continue
        parsed.append(item)
        if len(parsed) > MAX_ROWS:
            raise TienDoImportError(f'File vượt quá {MAX_ROWS} dòng.')

    if not parsed:
        raise TienDoImportError('File không có dòng dữ liệu nào.')
    return parsed


def import_rows(rows: list[ParsedRow], *, platform: str, user) -> int:
    """Tạo dòng tiến độ. Tạo ngược thứ tự để thứ tự hiển thị (mới nhất trước) khớp file."""
    objs = [
        TienDoItem(
            platform=platform,
            feature=r.feature,
            description=r.description,
            user_flow=r.user_flow,
            created_by=user,
        )
        for r in reversed(rows)
    ]
    for obj in objs:
        obj.save()
    return len(objs)


def html_to_plain(value: str) -> str:
    """Nội dung rich (text + ảnh) → chữ thuần cho Excel. Ảnh ghi [Ảnh]."""
    if not value:
        return ''
    text = re.sub(r'<img\b[^>]*>', '\n[Ảnh]\n', value, flags=re.I)
    text = re.sub(r'<br\s*/?>', '\n', text, flags=re.I)
    text = re.sub(r'</(p|div|li|ul|ol|h[1-6]|tr|table)>', '\n', text, flags=re.I)
    text = re.sub(r'<li\b[^>]*>', '- ', text, flags=re.I)
    text = re.sub(r'<[^>]+>', '', text)
    text = html.unescape(text).replace('\xa0', ' ')
    lines = [' '.join(line.split()) for line in text.replace('\r', '\n').split('\n')]
    return '\n'.join(line for line in lines if line).strip()


def build_board_export_xlsx(records: list[dict]) -> bytes:
    """Xuất các dòng tiến độ (mỗi feedback một dòng; dòng chưa có feedback vẫn xuất)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Tien do'
    ws.append(EXPORT_HEADERS)
    for record in records:
        ws.append((
            record.get('stt') or '',
            record.get('feature') or '',
            record.get('description') or '',
            record.get('user_flow') or '',
            record.get('author') or '',
            record.get('feedback_at') or '',
            record.get('feedback') or '',
            record.get('note') or '',
            record.get('updated_at') or '',
        ))

    header_fill = PatternFill('solid', fgColor='DC2626')
    for cell in ws[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = header_fill
        cell.alignment = Alignment(vertical='center', wrap_text=True)
    for row in ws.iter_rows(min_row=2, max_col=len(EXPORT_HEADERS)):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical='top')
    for col, width in zip('ABCDEFGHI', (6, 36, 48, 48, 22, 18, 40, 40, 18)):
        ws.column_dimensions[col].width = width
    last_row = max(ws.max_row, 1)
    ws.auto_filter.ref = f'A1:I{last_row}'
    ws.freeze_panes = 'A2'

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_template_xlsx() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Tien do'
    ws.append(TEMPLATE_HEADERS)
    ws.append((
        'Ví dụ: Đặt lịch xe công tác',
        'Cho phép nhân viên đặt lịch xe, quản lý duyệt.',
        '1. Vào menu Tiện ích → Đặt lịch\n2. Chọn ngày, nhập lý do\n3. Bấm Gửi',
    ))

    header_fill = PatternFill('solid', fgColor='DC2626')
    for cell in ws[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = header_fill
        cell.alignment = Alignment(vertical='center')
    for cell in ws[2]:
        cell.alignment = Alignment(wrap_text=True, vertical='top')
    for col, width in zip('ABC', (36, 50, 60)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = 'A2'

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
