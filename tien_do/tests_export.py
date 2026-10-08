import io

import openpyxl
from django.contrib.auth import get_user_model
from django.test import TestCase
from django.urls import reverse

from hrm.models import Profile
from tien_do.importing import EXPORT_HEADERS, build_board_export_xlsx, html_to_plain
from tien_do.models import TienDoFeedback, TienDoItem


class HtmlToPlainTests(TestCase):
    def test_strips_markup_and_marks_images(self):
        text = html_to_plain(
            '<p>D\u00f2ng 1<br>D\u00f2ng 2</p><img src="a.png" alt="anh">'
            '<ul><li>B\u01b0\u1edbc m\u1ed9t</li></ul>'
        )
        self.assertEqual(
            text,
            'D\u00f2ng 1\nD\u00f2ng 2\n[\u1ea2nh]\n- B\u01b0\u1edbc m\u1ed9t',
        )


class ExportWorkbookTests(TestCase):
    def test_writes_header_and_row(self):
        raw = build_board_export_xlsx([{
            'stt': 1,
            'feature': 'Dat lich',
            'description': 'Mo ta',
            'user_flow': '1. Vao',
            'author': 'An',
            'feedback_at': '08/10/2026 08:00',
            'feedback': 'On',
            'note': '',
            'updated_at': '08/10/2026 09:00',
        }])
        ws = openpyxl.load_workbook(io.BytesIO(raw)).active
        self.assertEqual([cell.value for cell in ws[1]], list(EXPORT_HEADERS))
        self.assertEqual(ws['B2'].value, 'Dat lich')
        self.assertEqual(ws['G2'].value, 'On')


class WholesaleRetailExportTests(TestCase):
    def setUp(self):
        user_model = get_user_model()
        self.user = user_model.objects.create_superuser('admin', 'td@example.com', 'pass-word-1')
        Profile.objects.filter(user=self.user).update(must_change_password=False)
        self.client.force_login(self.user)
        self.item = TienDoItem.objects.create(
            platform=TienDoItem.PLATFORM_WHOLESALE_RETAIL,
            feature='Gio hang',
            description='<p>Them san pham</p>',
            user_flow='1. Mo gio',
            created_by=self.user,
        )
        TienDoFeedback.objects.create(
            item=self.item,
            author=self.user,
            feedback='<p>On</p>',
            note='<p>Xem <img src="x.png"></p>',
        )
        TienDoFeedback.objects.create(
            item=self.item,
            author=self.user,
            feedback='Can sua nut',
            note='',
        )
        TienDoItem.objects.create(
            platform=TienDoItem.PLATFORM_WHOLESALE_RETAIL,
            feature='Chua co feedback',
            description='',
            user_flow='',
            created_by=self.user,
        )
        TienDoItem.objects.create(
            platform=TienDoItem.PLATFORM_PORTAL,
            feature='Khong xuat',
            created_by=self.user,
        )

    def _sheet(self, response):
        self.assertEqual(response.status_code, 200)
        self.assertIn('spreadsheetml', response['Content-Type'])
        return openpyxl.load_workbook(io.BytesIO(response.content)).active

    def test_export_includes_feedback_rows_and_skips_other_platform(self):
        response = self.client.get(reverse('tien_do:wholesale_retail_export'))
        self.assertIn('tien_do_website_si_le_', response['Content-Disposition'])
        ws = self._sheet(response)
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        features = [row[1] for row in rows]
        self.assertEqual(features.count('Gio hang'), 2)
        self.assertIn('Chua co feedback', features)
        self.assertNotIn('Khong xuat', features)
        gio_rows = [row for row in rows if row[1] == 'Gio hang']
        self.assertEqual(gio_rows[0][0], gio_rows[1][0])
        flat = ' '.join(str(cell or '') for row in rows for cell in row)
        self.assertIn('Them san pham', flat)
        self.assertIn('On', flat)
        self.assertIn('[' + chr(0x1ea2) + 'nh]', flat)
        self.assertNotIn('<p>', flat)

    def test_search_limits_exported_rows(self):
        response = self.client.get(reverse('tien_do:wholesale_retail_export'), {'q': 'Gio'})
        ws = self._sheet(response)
        features = [row[1] for row in ws.iter_rows(min_row=2, values_only=True)]
        self.assertEqual(set(features), {'Gio hang'})
