"""Tổng hợp chi phí NVL định mức theo nhóm NPL × mã hàng (hồ sơ đang dùng)."""

from __future__ import annotations

from dataclasses import dataclass, field
from decimal import Decimal
from io import BytesIO

from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from kho_npl.services.batches import material_avg_price
from san_xuat.models import BomVersion, ProductTechDoc
from san_xuat.services.costing import _d

ZERO = Decimal('0')
MONEY = Decimal('0.01')


@dataclass
class CostStatsProduct:
    doc_pk: int
    product_code: str
    product_name: str
    bom_pk: int | None
    bom_label: str
    bom_status: str = ''


@dataclass
class CostStatsCategory:
    category_id: int
    category_code: str
    category_name: str
    sort_order: int
    amounts: list[Decimal] = field(default_factory=list)

    @property
    def row_total(self) -> Decimal:
        return sum(self.amounts, ZERO)


@dataclass
class NvlCostMatrix:
    products: list[CostStatsProduct]
    categories: list[CostStatsCategory]
    column_totals: list[Decimal]
    grand_total: Decimal


def _column_label(doc: ProductTechDoc) -> str:
    name = (doc.product_name or '').strip()
    return name or (doc.product_code or '').strip() or f'#{doc.pk}'


def _unique_column_labels(docs: list[ProductTechDoc]) -> list[str]:
    """Ưu tiên product_name; nếu trùng tên thì thêm mã SP để phân biệt cột."""
    base = [_column_label(doc) for doc in docs]
    counts: dict[str, int] = {}
    for label in base:
        key = label.casefold()
        counts[key] = counts.get(key, 0) + 1
    labels: list[str] = []
    for doc, label in zip(docs, base):
        if counts[label.casefold()] > 1:
            code = (doc.product_code or '').strip()
            if code and code.casefold() != label.casefold():
                labels.append(f'{label} ({code})')
            else:
                labels.append(f'{label} #{doc.pk}')
        else:
            labels.append(label)
    return labels


def build_nvl_cost_matrix() -> NvlCostMatrix:
    """Ma trận: hàng = mọi nhóm NPL đang dùng, cột = mọi hồ sơ đang dùng."""
    from kho_npl.models import MaterialCategory
    from kho_npl.stock_domain import STOCK_DOMAIN_NPL
    from san_xuat.services.bom import get_working_bom

    docs = list(
        ProductTechDoc.objects.filter(is_active=True).order_by('product_code'),
    )
    bom_by_idx: list[BomVersion | None] = [get_working_bom(doc) for doc in docs]
    labels = _unique_column_labels(docs)
    products: list[CostStatsProduct] = [
        CostStatsProduct(
            doc_pk=doc.pk,
            product_code=doc.product_code or '',
            product_name=label,
            bom_pk=bom.pk if bom else None,
            bom_label=(bom.version_label or '') if bom else '',
            bom_status=bom.status if bom else '',
        )
        for doc, bom, label in zip(docs, bom_by_idx, labels)
    ]

    n_prod = len(products)
    amounts_by_cat: dict[int, list[Decimal]] = {}
    cat_meta: dict[int, tuple[str, str, int]] = {}

    for cat in (
        MaterialCategory.objects.filter(stock_domain=STOCK_DOMAIN_NPL, is_active=True)
        .order_by('sort_order', 'name', 'pk')
    ):
        amounts_by_cat[cat.pk] = [ZERO] * n_prod
        cat_meta[cat.pk] = (cat.code or '', cat.name or '', int(cat.sort_order or 0))

    for idx, bom in enumerate(bom_by_idx):
        if bom is None:
            continue
        lines = bom.lines.select_related(
            'material', 'material__category', 'material__unit',
        ).all()
        for line in lines:
            material = line.material
            cat = material.category
            cat_id = cat.pk
            if cat_id not in amounts_by_cat:
                amounts_by_cat[cat_id] = [ZERO] * n_prod
                cat_meta[cat_id] = (
                    cat.code or '',
                    cat.name or '',
                    int(cat.sort_order or 0),
                )
            unit_price = material_avg_price(material)
            amount = (line.qty_with_scrap * unit_price).quantize(MONEY)
            amounts_by_cat[cat_id][idx] += amount

    categories: list[CostStatsCategory] = [
        CostStatsCategory(
            category_id=cat_id,
            category_code=cat_meta[cat_id][0],
            category_name=cat_meta[cat_id][1],
            sort_order=cat_meta[cat_id][2],
            amounts=[_d(a).quantize(MONEY) for a in amounts],
        )
        for cat_id, amounts in amounts_by_cat.items()
    ]
    categories.sort(key=lambda c: (c.sort_order, c.category_name.casefold(), c.category_id))

    column_totals = [ZERO] * n_prod
    for cat in categories:
        for i, amount in enumerate(cat.amounts):
            column_totals[i] += amount
    column_totals = [a.quantize(MONEY) for a in column_totals]
    grand_total = sum(column_totals, ZERO).quantize(MONEY)

    return NvlCostMatrix(
        products=products,
        categories=categories,
        column_totals=column_totals,
        grand_total=grand_total,
    )


def export_nvl_cost_matrix_xlsx(matrix: NvlCostMatrix | None = None) -> HttpResponse:
    """Xuất Excel cùng bố cục mẫu th.xlsx (STT / Tên chi phí / mã hàng / Tổng)."""
    matrix = matrix or build_nvl_cost_matrix()
    products = matrix.products
    categories = matrix.categories
    n_prod = len(products)

    wb = Workbook()
    ws = wb.active
    ws.title = 'Tong_hop'

    thin = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1'),
    )
    header_fill = PatternFill('solid', fgColor='F1F5F9')
    section_fill = PatternFill('solid', fgColor='E2E8F0')
    total_fill = PatternFill('solid', fgColor='FEF3C7')
    title_font = Font(bold=True, size=14)
    bold = Font(bold=True)
    center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    left = Alignment(horizontal='left', vertical='center', wrap_text=True)
    right = Alignment(horizontal='right', vertical='center')

    last_col = 3 + n_prod  # A=STT B=Tên C..=products M=Tổng → col index
    total_col = last_col
    # Title
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(total_col, 2))
    ws['A1'] = 'Tổng hợp chi phí tính giá thành sản xuất'
    ws['A1'].font = title_font
    ws['A1'].alignment = Alignment(horizontal='center', vertical='center')

    # Header row 2–3
    ws.merge_cells('A2:A3')
    ws['A2'] = 'STT'
    ws['A2'].font = bold
    ws['A2'].fill = header_fill
    ws['A2'].alignment = center

    ws.merge_cells('B2:B3')
    ws['B2'] = 'Tên chi phí'
    ws['B2'].font = bold
    ws['B2'].fill = header_fill
    ws['B2'].alignment = center

    if n_prod:
        ws.merge_cells(start_row=2, start_column=3, end_row=2, end_column=2 + n_prod)
        ws.cell(2, 3, 'Mã hàng').font = bold
        ws.cell(2, 3).fill = header_fill
        ws.cell(2, 3).alignment = center
        for i, prod in enumerate(products):
            header = prod.product_name
            code = (prod.product_code or '').strip()
            if code and code not in header:
                header = f'{prod.product_name}\n{code}'
            cell = ws.cell(3, 3 + i, header)
            cell.font = bold
            cell.fill = header_fill
            cell.alignment = center
            cell.border = thin
    else:
        ws.cell(2, 3, 'Mã hàng').font = bold
        ws.cell(2, 3).fill = header_fill

    total_header = ws.cell(2, total_col, 'Tổng')
    ws.merge_cells(start_row=2, start_column=total_col, end_row=3, end_column=total_col)
    total_header.font = bold
    total_header.fill = header_fill
    total_header.alignment = center

    for col in range(1, total_col + 1):
        for row in (2, 3):
            c = ws.cell(row, col)
            c.border = thin
            if c.fill.fgColor is None or c.fill.fgColor.rgb in (None, '00000000'):
                c.fill = header_fill

    # Section I
    row = 4
    ws.cell(row, 1, 'I').font = bold
    ws.cell(row, 2, 'Chi phí NVL trực tiếp').font = bold
    for col in range(1, total_col + 1):
        c = ws.cell(row, col)
        c.fill = section_fill
        c.border = thin

    data_start = 5
    for i, cat in enumerate(categories, start=1):
        r = data_start + i - 1
        ws.cell(r, 1, i).alignment = center
        ws.cell(r, 2, cat.category_name).alignment = left
        for j, amount in enumerate(cat.amounts):
            cell = ws.cell(r, 3 + j, float(amount))
            cell.number_format = '#,##0.00'
            cell.alignment = right
        # Row total formula
        if n_prod:
            first = get_column_letter(3)
            last = get_column_letter(2 + n_prod)
            cell = ws.cell(r, total_col, f'=SUM({first}{r}:{last}{r})')
        else:
            cell = ws.cell(r, total_col, float(cat.row_total))
        cell.number_format = '#,##0.00'
        cell.alignment = right
        for col in range(1, total_col + 1):
            ws.cell(r, col).border = thin

    data_end = data_start + len(categories) - 1 if categories else data_start - 1
    cong_row = data_end + 1 if categories else data_start
    ws.cell(cong_row, 1, 'C').font = bold
    ws.cell(cong_row, 2, 'Cộng').font = bold
    for j in range(n_prod):
        col_letter = get_column_letter(3 + j)
        if categories:
            cell = ws.cell(
                cong_row, 3 + j,
                f'=SUM({col_letter}{data_start}:{col_letter}{data_end})',
            )
        else:
            cell = ws.cell(cong_row, 3 + j, 0)
        cell.number_format = '#,##0.00'
        cell.font = bold
        cell.fill = total_fill
        cell.alignment = right
    if n_prod:
        first = get_column_letter(3)
        last = get_column_letter(2 + n_prod)
        cell = ws.cell(cong_row, total_col, f'=SUM({first}{cong_row}:{last}{cong_row})')
    else:
        cell = ws.cell(cong_row, total_col, float(matrix.grand_total))
    cell.number_format = '#,##0.00'
    cell.font = bold
    cell.fill = total_fill
    for col in range(1, total_col + 1):
        c = ws.cell(cong_row, col)
        c.border = thin
        c.fill = total_fill

    ws.column_dimensions['A'].width = 6
    ws.column_dimensions['B'].width = 36
    for i in range(n_prod):
        ws.column_dimensions[get_column_letter(3 + i)].width = 14
    ws.column_dimensions[get_column_letter(total_col)].width = 14
    ws.row_dimensions[1].height = 24
    ws.freeze_panes = 'C4'

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    stamp = timezone.localtime().strftime('%Y%m%d_%H%M')
    response = HttpResponse(
        output.read(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = (
        f'attachment; filename=Thong_ke_chi_phi_NVL_{stamp}.xlsx'
    )
    return response
