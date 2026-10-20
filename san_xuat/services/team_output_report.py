"""Báo cáo sản lượng tổ theo ngày (tab Tổ SX trên bảng KHSX).

Hàng = tổ bộ phận HR, cột = ngày. SL ô = số bộ tổ làm xong trong ngày:
mỗi lệnh × size lấy min các công đoạn của tổ (như phiếu tiến độ tổ), rồi cộng
các size và các lệnh.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from datetime import date, timedelta
from decimal import Decimal
from types import SimpleNamespace

from django.db.models import Prefetch, Q

from san_xuat.hub_models import (
    SxMoProcessStep,
    SxOrderTeamDayPlan,
    SxProductionOrder,
    SxProductionStat,
)

ZERO = Decimal('0')

@dataclass
class TeamOutputCell:
    day: date
    qty: Decimal = ZERO
    plan: Decimal | None = None
    is_workday: bool = True
    is_today: bool = False
    is_past: bool = False
    orders: list[tuple[str, Decimal]] = field(default_factory=list)
    plan_orders: list[dict] = field(default_factory=list)
    actual_orders: list[dict] = field(default_factory=list)

    @property
    def has_value(self) -> bool:
        return bool(self.qty) or bool(self.plan)

    @property
    def state(self) -> str:
        """met / under / miss khi ngày đã qua có kế hoạch; '' nếu chưa tới hạn hoặc không KH."""
        if not self.plan or not self.is_past:
            return ''
        if self.qty >= self.plan:
            return 'met'
        return 'under' if self.qty > 0 else 'miss'

    @property
    def tip(self) -> str:
        from san_xuat.templatetags.sx_format import sx_num

        parts = [f'Dự kiến {sx_num(self.plan or 0)} · Thực hiện {sx_num(self.qty)}']
        parts += [f'{code}: {sx_num(qty)}' for code, qty in self.orders]
        return '\n'.join(parts)


@dataclass
class TeamOutputRow:
    slug: str
    label: str
    stage_slug: str
    stage_label: str
    color_slug: str = ''
    cells: list[TeamOutputCell] = field(default_factory=list)
    total: Decimal = ZERO
    plan_total: Decimal = ZERO
    active_days: int = 0
    avg_per_day: Decimal = ZERO
    best_day: date | None = None
    best_qty: Decimal = ZERO
    low_day: date | None = None
    low_qty: Decimal = ZERO
    order_count: int = 0
    stage_share: Decimal | None = None
    href: str = ''
    note: str = ''
    is_total: bool = False
    order_id: int = 0
    entry_mos: list = field(default_factory=list)
    entry_mos_json: str = '[]'
    closed_mos: list = field(default_factory=list)
    image_url: str = ''
    image_urls_json: str = '[]'
    priority: str = ''
    priority_label: str = ''
    progress_block: str = ''

    @property
    def plan_rate(self) -> Decimal | None:
        if self.plan_total <= 0:
            return None
        return (self.total * 100 / self.plan_total).quantize(Decimal('1'))

    @property
    def ready_to_complete(self) -> bool:
        """SL tổ đã vượt 80% KH — hiện nút Hoàn thành trên cột Đơn."""
        if self.is_total or self.plan_total <= 0:
            return False
        return (self.total * 100 / self.plan_total) > 80


@dataclass
class TeamOutputStage:
    slug: str
    label: str
    rows: list[TeamOutputRow] = field(default_factory=list)
    day_totals: list[Decimal] = field(default_factory=list)
    day_plans: list[Decimal | None] = field(default_factory=list)
    foot: list[TeamOutputCell] = field(default_factory=list)
    total: Decimal = ZERO
    plan_total: Decimal = ZERO
    active_days: int = 0
    avg_per_day: Decimal = ZERO

    @property
    def plan_rate(self) -> Decimal | None:
        if self.plan_total <= 0:
            return None
        return (self.total * 100 / self.plan_total).quantize(Decimal('1'))

    @property
    def best_qty(self) -> Decimal:
        return max(self.day_totals, default=ZERO)


@dataclass
class TeamOutputReport:
    date_from: date
    date_to: date
    days: list[TeamOutputCell]
    stages: list[TeamOutputStage]
    workday_count: int = 0
    active_team_count: int = 0
    idle_teams: list[str] = field(default_factory=list)
    unlabeled_qty: Decimal = ZERO
    day_details: dict = field(default_factory=dict)
    by_order: bool = False

    @property
    def rows(self) -> list[TeamOutputRow]:
        return [r for st in self.stages for r in st.rows]

    @property
    def has_data(self) -> bool:
        return any(st.total > 0 for st in self.stages)


def _stage_label(stage_slug: str) -> str:
    from san_xuat.services.progress_template import TEAM_SLUGS

    for slug, _gk, _mk, label in TEAM_SLUGS:
        if slug == stage_slug:
            return label
    return 'Chất lượng & khác'


def _q(v) -> Decimal:
    try:
        return Decimal(str(v or 0))
    except Exception:
        return ZERO


def _join_names(names) -> str:
    """Tên sản phẩm, không kèm mã."""
    cleaned: list[str] = []
    for name in names or []:
        text = ' '.join(str(name or '').split())
        if text and text not in cleaned:
            cleaned.append(text)
    if not cleaned:
        return ''
    label = ', '.join(cleaned[:3])
    if len(cleaned) > 3:
        label += '…'
    return label


def _priority_fields(order) -> dict:
    from san_xuat.hub_models import SxSalesOrder

    labels = dict(SxSalesOrder.PRIORITY_CHOICES)
    key = (getattr(order, 'plan_priority', '') or '').strip() or SxSalesOrder.PRIORITY_NORMAL
    if key not in labels:
        key = SxSalesOrder.PRIORITY_NORMAL
    return {'priority': key, 'priority_label': labels[key]}


def _team_day_qty_for_mo(
    mo: SxProductionOrder,
    *,
    sizes,
    keys: set[str],
    label_map: dict,
    stats: list[SxProductionStat],
    date_from: date,
    date_to: date,
) -> dict[date, Decimal]:
    """SL theo bộ của một tổ trên một lệnh, theo ngày.

    Dòng ghi không size (tổ nhập tổng) gom vào nhóm riêng `''` thay vì bỏ qua.
    """
    if not keys:
        return {}
    size_set = {r.size_label for r in sizes} | {''}
    single_total = len(sizes) == 1 and sizes[0].size_label == 'Tổng'
    acc: dict[date, dict[tuple[str, str], Decimal]] = {}
    for st in stats:
        if not (date_from <= st.stat_date <= date_to):
            continue
        step = label_map.get((st.process_name or '').strip().casefold())
        if not step or step.key not in keys:
            continue
        size = (st.size_label or '').strip()
        if single_total:
            size = 'Tổng'
        elif size not in size_set:
            continue
        qty = _q(st.qty_good)
        if qty <= 0:
            continue
        day = acc.setdefault(st.stat_date, {})
        day[(size, step.key)] = day.get((size, step.key), ZERO) + qty
    out: dict[date, Decimal] = {}
    for d, cells in acc.items():
        total = sum(
            (min((cells.get((sz, k), ZERO) for k in keys), default=ZERO) for sz in size_set),
            ZERO,
        )
        if total > 0:
            out[d] = total
    return out


def _day_plan_team_maps(teams):
    """Slug tổ HR, mã tổ năng lực, và các tổ cùng một công đoạn."""
    by_slug = {col.slug: col for col in teams}
    by_wc: dict[str, str] = {}
    teams_by_stage: dict[str, list[str]] = {}
    for col in teams:
        stage = (col.meta.get('stage_slug') or '').strip().lower()
        code = (col.meta.get('work_center_code') or '').strip().upper()
        if code:
            by_wc[code] = col.slug
        if stage:
            teams_by_stage.setdefault(stage, []).append(col.slug)
    return by_slug, by_wc, teams_by_stage


def _stage_claims_from_member_rows(rows, *, by_slug: dict, teams_by_stage: dict[str, list[str]]) -> dict[str, set[int]]:
    """Đơn đã có phân bổ ngày trên một tổ cụ thể của công đoạn.

    Slug cũ của công đoạn (``may``) không được cộng thêm vào tổ đó.
    """
    stage_of: dict[str, list[str]] = {}
    for stage, members in teams_by_stage.items():
        for member in members:
            stage_of.setdefault(member, []).append(stage)
    claimed: dict[str, set[int]] = {}
    for row in rows:
        key = (getattr(row, 'team_slug', '') or '').strip().lower()
        if key not in by_slug:
            continue
        oid = int(getattr(row, 'sales_order_id', 0) or 0)
        if not oid:
            continue
        for stage in stage_of.get(key, ()):
            claimed.setdefault(stage, set()).add(oid)
    return claimed


def _attribute_day_plan_team(
    team_slug: str,
    work_center_code: str,
    order_id: int,
    *,
    by_slug: dict,
    by_wc: dict[str, str],
    teams_by_stage: dict[str, list[str]],
    order_teams: dict[int, set[str]],
    stage_claims: dict[str, set[int]] | None = None,
) -> str:
    """Dòng ngày KHSX thuộc tổ HR nào.

    Mã `d{id}` hoặc mã tổ năng lực vào đúng bộ phận. Slug công đoạn cũ chỉ
    gom khi công đoạn còn một tổ, hoặc đơn chỉ được gán một tổ trong công đoạn.
    Đơn đã có ngày của tổ cụ thể thì bỏ slug cũ, không cộng đè.
    """
    key = (team_slug or '').strip().lower()
    if key in by_slug:
        return key
    wc = (work_center_code or '').strip().upper()
    if wc and wc in by_wc:
        return by_wc[wc]
    from san_xuat.services.capacity_from_hrm import stage_slug_is_shared

    # Slug chung (may) không còn là dòng trên lộ trình: lộ trình dùng d25 / d39.
    if stage_slug_is_shared(key):
        return ''
    if order_id in (stage_claims or {}).get(key, ()):
        return ''
    members = teams_by_stage.get(key) or []
    if len(members) == 1:
        return members[0]
    candidates = [slug for slug in members if slug in order_teams.get(order_id, ())]
    if len(candidates) == 1:
        return candidates[0]
    return ''


def retire_legacy_may_day_plans() -> dict[str, int]:
    """Bỏ slug ``may``. Đơn đã có tổ cụ thể thì xóa dòng cũ; đơn chỉ còn ``may`` thì chuyển sang đúng một tổ may."""
    from django.db import transaction
    from django.db.models import Q

    from san_xuat.hub_models import SxOrderTeamDayPlan, SxSalesOrderPlanStep
    from san_xuat.services.goods_progress import _team_columns

    _by_slug, by_wc, teams_by_stage = _day_plan_team_maps(_team_columns())
    members = set(teams_by_stage.get('may') or [])
    may_ids = set(
        SxOrderTeamDayPlan.objects.filter(team_slug__iexact='may')
        .values_list('sales_order_id', flat=True)
    )
    if not may_ids:
        return {'deleted': 0, 'moved': 0, 'left': 0}
    specific_ids = set()
    if members:
        specific_ids = set(
            SxOrderTeamDayPlan.objects.filter(
                sales_order_id__in=may_ids,
                team_slug__in=members,
            ).values_list('sales_order_id', flat=True)
        )
    deleted = 0
    moved = 0
    with transaction.atomic():
        if specific_ids:
            deleted = SxOrderTeamDayPlan.objects.filter(
                sales_order_id__in=specific_ids,
                team_slug__iexact='may',
            ).delete()[0]
        for oid in sorted(may_ids - specific_ids):
            teams = set()
            for code in SxSalesOrderPlanStep.objects.filter(
                sales_order_id=oid,
            ).values_list('work_center__code', flat=True):
                team = by_wc.get((code or '').strip().upper())
                if team in members:
                    teams.add(team)
            rows = SxOrderTeamDayPlan.objects.filter(sales_order_id=oid, team_slug__iexact='may')
            if len(teams) != 1:
                deleted += rows.delete()[0]
                continue
            target = next(iter(teams))
            for row in list(rows):
                clash = SxOrderTeamDayPlan.objects.filter(
                    sales_order_id=oid,
                    team_slug=target,
                    plan_date=row.plan_date,
                ).filter(
                    Q(work_center_id=row.work_center_id) if row.work_center_id else Q(work_center__isnull=True)
                ).exclude(pk=row.pk).first()
                if clash is not None:
                    clash.qty = _q(clash.qty) + _q(row.qty)
                    clash.save(update_fields=['qty'])
                    row.delete()
                    deleted += 1
                    continue
                row.team_slug = target
                row.save(update_fields=['team_slug'])
                moved += 1
    left = SxOrderTeamDayPlan.objects.filter(team_slug__iexact='may').count()
    return {'deleted': deleted, 'moved': moved, 'left': left}


def _plan_by_team_day(
    teams, *, date_from: date, date_to: date,
) -> tuple[
    dict[str, dict[date, Decimal]],
    dict[str, dict[date, Decimal]],
    dict[str, dict[date, list[dict]]],
]:
    """SL phân bổ ngày KHSX → (theo tổ, theo công đoạn, đơn của từng tổ/ngày)."""
    from django.urls import reverse

    from san_xuat.hub_models import SxSalesOrder, SxSalesOrderLine, SxSalesOrderPlanStep

    by_slug, by_wc, teams_by_stage = _day_plan_team_maps(teams)

    rows = list(
        SxOrderTeamDayPlan.objects.filter(
            plan_date__range=(date_from, date_to),
            sales_order__is_demo=False,
            qty__gt=0,
        ).select_related('sales_order', 'work_center')
    )
    order_ids = {row.sales_order_id for row in rows}
    products: dict[int, list[str]] = {}
    product_codes: dict[int, list[str]] = {}
    if order_ids:
        for oid, code, name in SxSalesOrderLine.objects.filter(
            order_id__in=order_ids,
        ).values_list('order_id', 'product_code', 'product_name'):
            code = (code or '').strip()
            if code and code not in product_codes.get(oid, ()):
                product_codes.setdefault(oid, []).append(code)
            label = (name or '').strip()
            if not label or label in products.get(oid, ()):
                continue
            products.setdefault(oid, []).append(label)
    from san_xuat.services.products import product_gallery_map

    galleries = product_gallery_map(
        code for codes in product_codes.values() for code in codes
    )

    order_teams: dict[int, set[str]] = {}
    if order_ids and by_wc:
        for oid, code in SxSalesOrderPlanStep.objects.filter(
            sales_order_id__in=order_ids,
            work_center__code__in=list(by_wc),
        ).values_list('sales_order_id', 'work_center__code'):
            team = by_wc.get((code or '').strip().upper())
            if team:
                order_teams.setdefault(oid, set()).add(team)

    by_team: dict[str, dict[date, Decimal]] = {}
    by_stage: dict[str, dict[date, Decimal]] = {}
    orders: dict[str, dict[date, dict[int, dict]]] = {}
    stage_claims = _stage_claims_from_member_rows(
        rows, by_slug=by_slug, teams_by_stage=teams_by_stage,
    )

    def _add(bucket: dict, key: str, day: date, qty: Decimal) -> None:
        inner = bucket.setdefault(key, {})
        inner[day] = inner.get(day, ZERO) + qty

    for row in rows:
        team = _attribute_day_plan_team(
            row.team_slug,
            getattr(row.work_center, 'code', '') or '',
            row.sales_order_id,
            by_slug=by_slug,
            by_wc=by_wc,
            teams_by_stage=teams_by_stage,
            order_teams=order_teams,
            stage_claims=stage_claims,
        )
        if not team:
            continue
        qty = _q(row.qty)
        day = row.plan_date
        col = by_slug[team]
        stage = (col.meta.get('stage_slug') or '').strip().lower()
        if stage:
            _add(by_stage, stage, day, qty)
        _add(by_team, team, day, qty)
        order = row.sales_order
        bucket = orders.setdefault(team, {}).setdefault(day, {})
        hit = bucket.get(order.pk)
        if hit is None:
            names = products.get(order.pk) or []
            product = _join_names(names)
            images: list[str] = []
            seen_images: set[str] = set()
            for code in product_codes.get(order.pk, []):
                for url in galleries.get(code.casefold()) or []:
                    if url and url not in seen_images:
                        seen_images.add(url)
                        images.append(url)
            bucket[order.pk] = {
                'id': order.pk,
                'code': order.code,
                **_priority_fields(order),
                'customer': (order.customer_name or '').strip(),
                'product': product,
                'image': images[0] if images else '',
                'images': images[:6],
                'qty': qty,
                'due': order.due_date.strftime('%d/%m/%Y') if order.due_date else '',
                'url': reverse('san_xuat:sales_order_detail', args=[order.pk]),
            }
        else:
            hit['qty'] = hit['qty'] + qty
    covered = {
        (team, int(oid))
        for team, days in orders.items()
        for day_items in days.values()
        for oid in day_items
    }
    extras = [
        item for item in _route_span_fallbacks(teams, covered=covered)
        if date_from <= item[1] <= date_to
    ]
    if extras:
        missing = [oid for _team, _day, oid, _status, _qty in extras if oid not in products]
        if missing:
            for oid, code, name in SxSalesOrderLine.objects.filter(
                order_id__in=missing,
            ).values_list('order_id', 'product_code', 'product_name'):
                code = (code or '').strip()
                if code and code not in product_codes.get(oid, ()):
                    product_codes.setdefault(oid, []).append(code)
                label = (name or '').strip()
                if label and label not in products.get(oid, ()):
                    products.setdefault(oid, []).append(label)
            galleries.update(product_gallery_map(
                code for oid in missing for code in product_codes.get(oid, [])
            ))
        extra_orders = {
            order.pk: order
            for order in SxSalesOrder.objects.filter(pk__in=[oid for _t, _d, oid, _s, _q in extras])
        }
        for team, day, oid, _status, qty in extras:
            order = extra_orders.get(oid)
            if order is None:
                continue
            _add(by_team, team, day, qty)
            col = by_slug.get(team)
            stage = (col.meta.get('stage_slug') or '').strip().lower() if col is not None else ''
            if stage:
                _add(by_stage, stage, day, qty)
            bucket = orders.setdefault(team, {}).setdefault(day, {})
            if oid in bucket:
                bucket[oid]['qty'] = bucket[oid]['qty'] + qty
                continue
            names = products.get(oid) or []
            images = []
            seen_images: set[str] = set()
            for code in product_codes.get(oid, []):
                for url in galleries.get(code.casefold()) or []:
                    if url and url not in seen_images:
                        seen_images.add(url)
                        images.append(url)
            bucket[oid] = {
                'id': oid,
                'code': order.code,
                **_priority_fields(order),
                'customer': (order.customer_name or '').strip(),
                'product': _join_names(names),
                'image': images[0] if images else '',
                'images': images[:6],
                'qty': qty,
                'due': order.due_date.strftime('%d/%m/%Y') if order.due_date else '',
                'url': reverse('san_xuat:sales_order_detail', args=[oid]),
            }
    listed = {
        team: {day: list(items.values()) for day, items in days.items()}
        for team, days in orders.items()
    }
    return by_team, by_stage, listed


def build_team_output_report(
    *,
    date_from: date,
    date_to: date,
    search: str = '',
    today: date | None = None,
) -> TeamOutputReport:
    from django.urls import reverse
    from django.utils import timezone

    from san_xuat.services.goods_progress import _team_columns
    from san_xuat.services.order_progress_sheet import (
        _size_plans,
        progress_steps_for_mo,
        progress_steps_for_team,
        stats_for_team,
    )
    from san_xuat.services.work_calendar import working_days

    today = today or timezone.localdate()
    from san_xuat.services.plan_board import _factory_slug_rank

    rank = _factory_slug_rank()
    teams = sorted(
        _team_columns(),
        key=lambda col: (rank.get(col.slug, 99), col.label),
    )
    workdays = set(working_days(date_from, date_to))
    day_list: list[date] = []
    d = date_from
    while d <= date_to:
        day_list.append(d)
        d += timedelta(days=1)

    stats_qs = SxProductionStat.objects.filter(
        is_demo=False,
        status=SxProductionStat.STATUS_CONFIRMED,
        stat_date__range=(date_from, date_to),
        production_order__isnull=False,
    )
    if search:
        stats_qs = stats_qs.filter(
            Q(production_order__code__icontains=search)
            | Q(production_order__product_code__icontains=search)
            | Q(production_order__product_name__icontains=search)
            | Q(production_order__sales_order__code__icontains=search)
        )
    stats_by_mo: dict[int, list[SxProductionStat]] = {}
    unlabeled = ZERO
    for st in stats_qs.only(
        'production_order_id', 'stat_date', 'process_name', 'size_label', 'qty_good', 'team_label',
    ):
        if not (st.team_label or '').strip():
            unlabeled += _q(st.qty_good)
            continue
        stats_by_mo.setdefault(st.production_order_id, []).append(st)

    mos = list(
        SxProductionOrder.objects.filter(pk__in=list(stats_by_mo))
        .select_related('sales_order')
        .prefetch_related(
            Prefetch(
                'mo_process_steps',
                queryset=SxMoProcessStep.objects.select_related('work_center').order_by('sequence', 'id'),
            ),
        )
    )

    qty: dict[str, dict[date, Decimal]] = {}
    orders: dict[str, dict[date, list[tuple[str, Decimal]]]] = {}
    actuals: dict[str, dict[date, dict[str, dict]]] = {}
    team_mos: dict[str, set[int]] = {}
    for mo in mos:
        mo_steps = progress_steps_for_mo(mo)
        if not mo_steps:
            continue
        sizes = _size_plans(mo)
        if not sizes:
            continue
        label_map = {s.label.casefold(): s for s in mo_steps}
        mo_stats = stats_by_mo.get(mo.pk, [])
        for col in teams:
            steps = progress_steps_for_team(mo, col.meta, steps=mo_steps)
            if not steps:
                continue
            per_day = _team_day_qty_for_mo(
                mo,
                sizes=sizes,
                keys={s.key for s in steps},
                label_map=label_map,
                stats=stats_for_team(mo_stats, col.meta),
                date_from=date_from,
                date_to=date_to,
            )
            if not per_day:
                continue
            team_mos.setdefault(col.slug, set()).add(mo.pk)
            bucket = qty.setdefault(col.slug, {})
            ob = orders.setdefault(col.slug, {})
            so = mo.sales_order
            so_code = (so.code if so else '').strip()
            for day, v in per_day.items():
                bucket[day] = bucket.get(day, ZERO) + v
                ob.setdefault(day, []).append((mo.code, v))
                if not so_code:
                    continue
                day_actual = actuals.setdefault(col.slug, {}).setdefault(day, {})
                hit = day_actual.get(so_code)
                if hit is None:
                    names = []
                    extra = (mo.product_name or '').strip()
                    if extra:
                        names.append(extra)
                    day_actual[so_code] = {
                        'id': so.pk,
                        'code': so_code,
                        **_priority_fields(so),
                        'customer': (so.customer_name or '').strip(),
                        'names': names,
                        'product': _join_names(names),
                        'product_code': (mo.product_code or '').strip(),
                        'qty': v,
                        'due': so.due_date.strftime('%d/%m/%Y') if so.due_date else '',
                        'url': reverse('san_xuat:sales_order_detail', args=[so.pk]),
                    }
                else:
                    hit['qty'] = hit['qty'] + v
                    extra = (mo.product_name or '').strip()
                    names = hit.setdefault('names', [])
                    if extra and extra not in names:
                        names.append(extra)
                        hit['product'] = _join_names(names)

    plans, stage_plans, plan_orders = _plan_by_team_day(teams, date_from=date_from, date_to=date_to)
    from san_xuat.services.team_stage_colors import team_stage_palette

    palette = team_stage_palette()

    stages: dict[str, TeamOutputStage] = {}
    for col in teams:
        stage_slug = (col.meta.get('stage_slug') or '').strip().lower()
        stage = stages.get(stage_slug)
        if stage is None:
            stage = TeamOutputStage(
                slug=stage_slug,
                label=_stage_label(stage_slug),
                day_totals=[ZERO] * len(day_list),
            )
            stages[stage_slug] = stage
        row = TeamOutputRow(
            slug=col.slug,
            label=col.label,
            stage_slug=stage_slug,
            stage_label=stage.label,
            color_slug=col.slug if col.slug in palette else (stage_slug or col.slug),
            order_count=len(team_mos.get(col.slug, ())),
        )
        t_qty = qty.get(col.slug, {})
        t_plan = plans.get(col.slug, {})
        t_orders = orders.get(col.slug, {})
        t_plan_orders = plan_orders.get(col.slug, {})
        for i, day in enumerate(day_list):
            v = t_qty.get(day, ZERO)
            cell = TeamOutputCell(
                day=day,
                qty=v,
                plan=t_plan.get(day),
                is_workday=day in workdays,
                is_today=day == today,
                is_past=day <= today,
                orders=sorted(t_orders.get(day, []), key=lambda t: -t[1]),
                plan_orders=sorted(t_plan_orders.get(day, []), key=lambda item: item['code']),
                actual_orders=list(actuals.get(col.slug, {}).get(day, {}).values()),
            )
            row.cells.append(cell)
            stage.day_totals[i] += v
            row.total += v
            row.plan_total += t_plan.get(day) or ZERO
        active = [c for c in row.cells if c.qty > 0]
        row.active_days = len(active)
        if active:
            row.avg_per_day = (row.total / len(active)).quantize(Decimal('1'))
            best = max(active, key=lambda c: c.qty)
            low = min(active, key=lambda c: c.qty)
            row.best_day, row.best_qty = best.day, best.qty
            row.low_day, row.low_qty = low.day, low.qty
        stage.rows.append(row)
        stage.total += row.total

    for stage in stages.values():
        s_plan = stage_plans.get(stage.slug, {})
        stage.day_plans = [s_plan.get(day) for day in day_list]
        stage.plan_total = sum(s_plan.values(), ZERO)
        stage.foot = [
            TeamOutputCell(
                day=day,
                qty=stage.day_totals[i],
                plan=stage.day_plans[i],
                is_workday=day in workdays,
                is_today=day == today,
                is_past=day <= today,
            )
            for i, day in enumerate(day_list)
        ]

    stage_order = [s.slug for s in teams_stage_order(teams)]
    ordered = sorted(stages.values(), key=lambda s: stage_order.index(s.slug) if s.slug in stage_order else 99)
    for stage in ordered:
        active_days = sum(1 for v in stage.day_totals if v > 0)
        stage.active_days = active_days
        if active_days:
            stage.avg_per_day = (stage.total / active_days).quantize(Decimal('1'))
        for row in stage.rows:
            if stage.total > 0 and len(stage.rows) > 1:
                row.stage_share = (row.total * 100 / stage.total).quantize(Decimal('1'))

    report = TeamOutputReport(
        date_from=date_from,
        date_to=date_to,
        days=[
            TeamOutputCell(day=d, is_workday=d in workdays, is_today=d == today)
            for d in day_list
        ],
        stages=ordered,
        workday_count=len(workdays),
        unlabeled_qty=unlabeled,
    )
    rows = report.rows
    report.active_team_count = sum(1 for r in rows if r.total > 0)
    report.idle_teams = [r.label for r in rows if r.total <= 0]
    report.day_details = _day_detail_payload(rows)
    return report


def _route_span_fallbacks(teams, *, covered: set[tuple[str, int]]):
    """Đơn lộ trình chưa có dòng ngày đúng slug bộ phận.

    Lộ trình khi đó vẽ một ô ở ngày bắt đầu, đủ sản lượng. Các dòng slug chung
    cũ (``may``) không được trải ra từng ngày trên lưới tổ.
    """
    from san_xuat.hub_models import SxSalesOrder, SxSalesOrderPlanStep
    from san_xuat.services.capacity_from_hrm import stage_slug_is_shared
    from san_xuat.services.plan_board import team_khsx_spans

    code_to_slug: dict[str, str] = {}
    shared_codes: set[str] = set()
    span_teams: dict[str, set[str]] = {}
    for col in teams:
        stage = (col.meta.get('stage_slug') or '').strip().lower()
        code = (col.meta.get('work_center_code') or '').strip().upper()
        if not code:
            continue
        code_to_slug[code] = col.slug
        span_teams.setdefault(col.slug, set()).add(col.slug)
        if stage:
            span_teams.setdefault(stage, set()).add(col.slug)
        if stage_slug_is_shared(stage):
            shared_codes.add(code)
    if not code_to_slug:
        return []
    need: dict[int, set[str]] = {}
    status_of: dict[int, str] = {}
    for oid, code, plan_status in SxSalesOrderPlanStep.objects.filter(
        work_center__code__in=list(code_to_slug),
        sales_order__is_demo=False,
    ).values_list('sales_order_id', 'work_center__code', 'sales_order__plan_status'):
        team = code_to_slug.get((code or '').strip().upper())
        if not team or (team, int(oid)) in covered:
            continue
        cancelled = plan_status == SxSalesOrder.PLAN_CANCELLED
        if not cancelled and (code or '').strip().upper() not in shared_codes:
            continue
        need.setdefault(int(oid), set()).add(team)
        if cancelled:
            status_of[int(oid)] = 'cancelled'
        elif plan_status == SxSalesOrder.PLAN_DONE:
            status_of[int(oid)] = 'done'
        else:
            status_of[int(oid)] = 'open'
    if not need:
        return []
    out: list[tuple[str, date, int, str, Decimal]] = []
    orders = SxSalesOrder.objects.filter(pk__in=list(need)).prefetch_related(
        'lines', 'plan_steps__work_center', 'team_day_plans',
    )
    for order in orders:
        want = need.get(order.pk) or set()
        for span in team_khsx_spans(order):
            slug = (span.slug or '').strip().lower()
            targets = (span_teams.get(slug) or set()) & want
            if not targets or span.start is None:
                continue
            qty = _q(getattr(span, 'planned_qty', 0) or 0)
            if qty <= 0:
                continue
            status = status_of.get(order.pk, 'open')
            for team in targets:
                out.append((team, span.start, int(order.pk), status, qty))
    return out


def _team_closed_sales_order_ids(team_slug: str) -> set[int]:
    """Đơn mà tổ này đã bấm Hoàn thành — chuyển sang tab Đã hoàn thành."""
    from san_xuat.hub_models import SxProductionOrder, SxTeamWorkClose
    from san_xuat.services.team_division_map import team_slug_aliases

    slug = (team_slug or '').strip().lower()
    if not slug:
        return set()
    mo_ids = SxTeamWorkClose.objects.filter(
        team_slug__in=team_slug_aliases(slug),
        is_demo=False,
    ).values_list('production_order_id', flat=True)
    return set(
        SxProductionOrder.objects.filter(
            pk__in=mo_ids,
            is_demo=False,
            sales_order_id__isnull=False,
        ).exclude(
            status=SxProductionOrder.STATUS_CANCELLED,
        ).values_list('sales_order_id', flat=True)
    )


def _team_plan_hits(team_slug: str) -> list[tuple[date, int, str]]:
    """(ngày, id đơn, trạng thái) của đơn được gán cho đúng tổ này.

    Trạng thái là ``done``, ``cancelled`` hoặc ``open``.
    """
    from san_xuat.hub_models import SxOrderTeamDayPlan, SxSalesOrder, SxSalesOrderPlanStep
    from san_xuat.services.goods_progress import _team_columns

    slug = (team_slug or '').strip().lower()
    by_slug, by_wc, teams_by_stage = _day_plan_team_maps(_team_columns())
    rows = list(
        SxOrderTeamDayPlan.objects.filter(
            qty__gt=0,
            sales_order__is_demo=False,
        ).select_related('work_center', 'sales_order')
    )
    order_ids = {row.sales_order_id for row in rows}
    order_teams: dict[int, set[str]] = {}
    if order_ids and by_wc:
        for oid, code in SxSalesOrderPlanStep.objects.filter(
            sales_order_id__in=order_ids,
            work_center__code__in=list(by_wc),
        ).values_list('sales_order_id', 'work_center__code'):
            team = by_wc.get((code or '').strip().upper())
            if team:
                order_teams.setdefault(oid, set()).add(team)
    stage_claims = _stage_claims_from_member_rows(
        rows, by_slug=by_slug, teams_by_stage=teams_by_stage,
    )
    hits: list[tuple[date, int, str]] = []
    for row in rows:
        if _attribute_day_plan_team(
            row.team_slug,
            getattr(row.work_center, 'code', '') or '',
            row.sales_order_id,
            by_slug=by_slug,
            by_wc=by_wc,
            teams_by_stage=teams_by_stage,
            order_teams=order_teams,
            stage_claims=stage_claims,
        ) != slug:
            continue
        if row.sales_order.plan_status == SxSalesOrder.PLAN_CANCELLED:
            status = 'cancelled'
        elif row.sales_order.plan_status == SxSalesOrder.PLAN_DONE:
            status = 'done'
        else:
            status = 'open'
        hits.append((row.plan_date, int(row.sales_order_id), status))
    covered = {(slug, oid) for _day, oid, _status in hits}
    for team, day, oid, status, _qty in _route_span_fallbacks(_team_columns(), covered=covered):
        if team == slug:
            hits.append((day, oid, status))
    closed_orders = _team_closed_sales_order_ids(slug)
    if not closed_orders:
        return hits
    return [
        (day, oid, 'done' if status == 'open' and oid in closed_orders else status)
        for day, oid, status in hits
    ]


def unfinished_sheet_bounds(today: date, team_slug: str) -> tuple[date, date]:
    """Lưới của một tổ: từ ngày sớm nhất của đơn còn lại đến hôm nay.

    Đơn đã hoàn thành không kéo ngày bắt đầu. Tổ cắt xong đến 10/10 thì lưới
    tổ cắt bắt đầu 11/10, trong khi tổ khác vẫn theo đơn còn của tổ đó.
    """
    days = [day for day, _oid, status in _team_plan_hits(team_slug) if status == 'open']
    return _sheet_bounds(days, today, extend_to_today=True)


def team_sheet_tab(
    today: date,
    team_slug: str,
    *,
    status: str,
    search: str = '',
    priority: str = '',
    deadline: str = '',
) -> tuple[date, date, int, int, int, set[int]]:
    """Cửa sổ lưới, số đơn ba tab, và id đơn của tab đang mở."""
    hits = _team_plan_hits(team_slug)
    open_ids = _filter_sheet_order_ids(
        {oid for _day, oid, hit_status in hits if hit_status == 'open'},
        search,
        priority=priority,
        deadline=deadline,
        today=today,
    )
    done_ids = _filter_sheet_order_ids(
        {oid for _day, oid, hit_status in hits if hit_status == 'done'},
        search,
        priority=priority,
        deadline=deadline,
        today=today,
    )
    cancelled_ids = _filter_sheet_order_ids(
        {oid for _day, oid, hit_status in hits if hit_status == 'cancelled'},
        search,
        priority=priority,
        deadline=deadline,
        today=today,
    )
    groups = {'open': open_ids, 'done': done_ids, 'cancelled': cancelled_ids}
    active = status if status in groups else 'open'
    keep = groups[active]
    days = [day for day, oid, hit_status in hits if hit_status == active and oid in keep]
    if active == 'open':
        start, end = _sheet_bounds(days, today, extend_to_today=True)
    else:
        start, end = _history_sheet_bounds(days, today)
    return start, end, len(open_ids), len(done_ids), len(cancelled_ids), keep


def _history_sheet_bounds(days: list[date], today: date) -> tuple[date, date]:
    """Tab hoàn thành và đã hủy luôn có ít nhất 30 ngày, kể cả khi chưa có đơn."""
    floor = today - timedelta(days=29)
    if not days:
        return floor, today
    start, end = min(days), max(days)
    if start > floor:
        start = floor
    if end < today:
        end = today
    return start, end


def _sheet_bounds(days: list[date], today: date, *, extend_to_today: bool) -> tuple[date, date]:
    if not days:
        return today, today
    start, end = min(days), max(days)
    if extend_to_today and end < today:
        end = today
    return start, end


def _filter_sheet_order_ids(
    ids: set[int],
    search: str,
    *,
    priority: str = '',
    deadline: str = '',
    today: date | None = None,
) -> set[int]:
    query = (search or '').strip()
    priority = (priority or '').strip()
    deadline = (deadline or '').strip()
    if not ids:
        return set()
    if not query and not priority and deadline not in ('overdue', 'due_soon', 'on_time'):
        return set(ids)
    from django.db.models import Q

    from san_xuat.hub_models import SxSalesOrder

    qs = SxSalesOrder.objects.filter(pk__in=ids)
    if query:
        qs = qs.filter(
            Q(code__icontains=query)
            | Q(customer_name__icontains=query)
            | Q(lines__product_code__icontains=query)
            | Q(lines__product_name__icontains=query)
            | Q(production_orders__code__icontains=query)
            | Q(production_orders__product_code__icontains=query)
            | Q(production_orders__product_name__icontains=query)
        )
    valid_priorities = {value for value, _label in SxSalesOrder.PRIORITY_CHOICES}
    if priority in valid_priorities:
        qs = qs.filter(plan_priority=priority)
    if today and deadline == 'overdue':
        qs = qs.filter(due_date__lt=today)
    elif today and deadline == 'due_soon':
        qs = qs.filter(due_date__gte=today, due_date__lte=today + timedelta(days=3))
    elif today and deadline == 'on_time':
        qs = qs.filter(Q(due_date__isnull=True) | Q(due_date__gte=today))
    return set(qs.values_list('pk', flat=True).distinct())


def _recompute_output_row(row: TeamOutputRow) -> None:
    row.total = ZERO
    row.plan_total = ZERO
    for cell in row.cells:
        row.total += cell.qty
        row.plan_total += cell.plan or ZERO
    active = [cell for cell in row.cells if cell.qty > 0]
    row.active_days = len(active)
    row.avg_per_day = (row.total / len(active)).quantize(Decimal('1')) if active else ZERO
    if active:
        best = max(active, key=lambda cell: cell.qty)
        row.best_day, row.best_qty = best.day, best.qty
    else:
        row.best_day, row.best_qty = None, ZERO


def retain_unfinished_orders(report: TeamOutputReport) -> TeamOutputReport:
    """Bỏ đơn đã hoàn thành hoặc đã hủy khỏi lưới và modal ngày."""
    from san_xuat.hub_models import SxSalesOrder

    seen: set[int] = set()
    for row in report.rows:
        for cell in row.cells:
            for item in (*cell.plan_orders, *cell.actual_orders):
                try:
                    oid = int(item.get('id') or 0)
                except (TypeError, ValueError):
                    oid = 0
                if oid:
                    seen.add(oid)
    open_ids = set(
        SxSalesOrder.objects.filter(pk__in=seen, is_demo=False)
        .exclude(plan_status__in=(SxSalesOrder.PLAN_DONE, SxSalesOrder.PLAN_CANCELLED))
        .values_list('pk', flat=True)
    ) if seen else set()
    return retain_order_ids(report, open_ids)


def retain_order_ids(report: TeamOutputReport, keep_ids: set[int]) -> TeamOutputReport:
    """Giữ đúng các đơn của tab đang mở. Cộng và modal ngày tính lại theo nhóm đó."""
    for row in report.rows:
        for cell in row.cells:
            cell.plan_orders = [
                item for item in cell.plan_orders if int(item.get('id') or 0) in keep_ids
            ]
            cell.actual_orders = [
                item for item in cell.actual_orders if int(item.get('id') or 0) in keep_ids
            ]
            plan = sum((_q(item.get('qty')) for item in cell.plan_orders), ZERO)
            cell.plan = plan if cell.plan_orders else None
            cell.qty = sum((_q(item.get('qty')) for item in cell.actual_orders), ZERO)
            if not cell.plan_orders and not cell.actual_orders:
                cell.orders = []
        _recompute_output_row(row)
    report.day_details = _day_detail_payload(report.rows)
    return report


def restrict_team_output_report(report: TeamOutputReport, team_slug: str) -> TeamOutputReport:
    """Giữ một tổ trên bản sao logic của tab Tổ SX — modal ngày không còn tổ khác."""
    slug = (team_slug or '').strip().lower()
    kept: list[TeamOutputStage] = []
    for stage in report.stages:
        rows = [row for row in stage.rows if (row.slug or '').strip().lower() == slug]
        if not rows:
            continue
        stage.rows = rows
        kept.append(stage)
    report.stages = kept
    prefix = f'{slug}|'
    report.day_details = {
        key: value
        for key, value in (report.day_details or {}).items()
        if str(key).startswith(prefix)
    }
    report.active_team_count = sum(1 for row in report.rows if row.total > 0)
    report.idle_teams = [row.label for row in report.rows if row.total <= 0]
    return report


def _order_rows_from_team(source: TeamOutputRow) -> list[TeamOutputRow]:
    """Tách một dòng tổ thành từng đơn. Ô ngày chỉ còn SL của đơn đó."""
    found: dict[str, dict] = {}

    def slot(item: dict) -> dict | None:
        try:
            oid = int(item.get('id') or 0)
        except (TypeError, ValueError):
            oid = 0
        code = (item.get('code') or '').strip()
        key = f'id:{oid}' if oid else f'code:{code}'
        if key in ('id:0', 'code:'):
            return None
        hit = found.get(key)
        if hit is None:
            hit = {
                'id': oid,
                'code': code or '—',
                'product': (item.get('product') or '').strip(),
                'product_code': (item.get('product_code') or '').strip(),
                'url': item.get('url') or '',
                'priority_key': (item.get('priority') or '').strip(),
                'priority_label': (item.get('priority_label') or '').strip(),
                'image': (item.get('image') or '').strip(),
                'images': list(item.get('images') or []),
                'plan': {},
                'done': {},
            }
            found[key] = hit
        else:
            if not hit['product'] and (item.get('product') or '').strip():
                hit['product'] = (item.get('product') or '').strip()
            if not hit['product_code'] and (item.get('product_code') or '').strip():
                hit['product_code'] = (item.get('product_code') or '').strip()
            if not hit['image'] and (item.get('image') or '').strip():
                hit['image'] = (item.get('image') or '').strip()
                hit['images'] = list(item.get('images') or []) or [hit['image']]
            if not hit['priority_key'] and (item.get('priority') or '').strip():
                hit['priority_key'] = (item.get('priority') or '').strip()
                hit['priority_label'] = (item.get('priority_label') or '').strip()
        return hit

    for cell in source.cells:
        for item in cell.plan_orders:
            hit = slot(item)
            if hit is None:
                continue
            hit['plan'][cell.day] = hit['plan'].get(cell.day, ZERO) + _q(item.get('qty'))
        for item in cell.actual_orders:
            hit = slot(item)
            if hit is None:
                continue
            hit['done'][cell.day] = hit['done'].get(cell.day, ZERO) + _q(item.get('qty'))

    missing_codes = [
        hit['product_code'] for hit in found.values() if not hit.get('image') and hit.get('product_code')
    ]
    if missing_codes:
        from san_xuat.services.products import product_gallery_map

        galleries = product_gallery_map(missing_codes)
        for hit in found.values():
            if hit.get('image'):
                continue
            urls = galleries.get((hit.get('product_code') or '').casefold()) or []
            if urls:
                hit['image'] = urls[0]
                hit['images'] = urls[:6]

    import json

    rows: list[TeamOutputRow] = []
    ordered = sorted(
        found.values(),
        key=lambda hit: (min(hit['plan'], default=date.max), hit['code']),
    )
    for hit in ordered:
        cells: list[TeamOutputCell] = []
        total = ZERO
        plan_total = ZERO
        for src in source.cells:
            plan = hit['plan'].get(src.day)
            done = hit['done'].get(src.day, ZERO)
            if plan is not None:
                plan_total += plan
            total += done
            cells.append(TeamOutputCell(
                day=src.day,
                qty=done,
                plan=plan,
                is_workday=src.is_workday,
                is_today=src.is_today,
                is_past=src.is_past,
            ))
        row = TeamOutputRow(
            slug=source.slug,
            label=hit['code'],
            stage_slug=source.stage_slug,
            stage_label=source.stage_label,
            color_slug=source.color_slug,
            cells=cells,
            total=total,
            plan_total=plan_total,
            href=hit['url'],
            order_id=int(hit.get('id') or 0),
            image_url=hit.get('image') or '',
            image_urls_json=json.dumps(hit.get('images') or ([hit['image']] if hit.get('image') else []), ensure_ascii=False),
            note=hit['product'],
            priority=hit.get('priority_key') or '',
            priority_label=hit.get('priority_label') or '',
        )
        active = [c for c in cells if c.qty > 0]
        row.active_days = len(active)
        if active:
            row.avg_per_day = (row.total / len(active)).quantize(Decimal('1'))
            best = max(active, key=lambda c: c.qty)
            row.best_day = best.day
            row.best_qty = best.qty
        rows.append(row)
    return rows


def attach_sheet_entry_mos(report: TeamOutputReport, team: dict) -> TeamOutputReport:
    """Gắn lệnh đang mở của tổ vào từng dòng đơn, để ô ngày mở phiếu nhập SL."""
    import json

    from san_xuat.hub_models import SxProductionOrder
    from san_xuat.services.order_progress_sheet import progress_steps_for_mo, progress_steps_for_team
    from san_xuat.services.team_work import team_job_closes

    order_ids = [row.order_id for row in report.rows if row.order_id and not row.is_total]
    if not order_ids:
        return report
    all_mos = list(
        SxProductionOrder.objects.filter(
            sales_order_id__in=order_ids,
            is_demo=False,
        )
        .select_related('sales_order')
        .prefetch_related('mo_process_steps__work_center')
        .order_by('-order_date', '-pk')
    )
    usable = [
        mo for mo in all_mos
        if mo.status not in (
            SxProductionOrder.STATUS_DRAFT,
            SxProductionOrder.STATUS_CANCELLED,
        )
    ]
    by_order: dict[int, list] = {}
    any_by_order: dict[int, list] = {}

    def _fill_maps(mos: list) -> None:
        any_by_order.clear()
        by_order.clear()
        for mo in mos:
            if mo.sales_order_id:
                any_by_order.setdefault(int(mo.sales_order_id), []).append(mo)
        for mo in mos:
            if mo.status in (
                SxProductionOrder.STATUS_DRAFT,
                SxProductionOrder.STATUS_CANCELLED,
            ):
                continue
            if not mo.sales_order_id:
                continue
            steps = progress_steps_for_team(mo, team, steps=progress_steps_for_mo(mo))
            if not steps:
                continue
            by_order.setdefault(int(mo.sales_order_id), []).append(mo)

    _fill_maps(all_mos)
    need_ids = []
    for row in report.rows:
        if row.is_total or not row.order_id or by_order.get(row.order_id):
            continue
        owned = any_by_order.get(row.order_id) or []
        if any(
            mo.status not in (
                SxProductionOrder.STATUS_DRAFT,
                SxProductionOrder.STATUS_CANCELLED,
            )
            for mo in owned
        ):
            continue
        need_ids.append(row.order_id)
    if need_ids:
        from san_xuat.hub_models import SxSalesOrder
        from san_xuat.services.order_progress_sheet import ensure_order_progress_mos

        for order in SxSalesOrder.objects.filter(pk__in=need_ids, is_demo=False):
            ensure_order_progress_mos(order)
        opened = list(
            SxProductionOrder.objects.filter(
                sales_order_id__in=need_ids,
                is_demo=False,
            )
            .select_related('sales_order')
            .prefetch_related('mo_process_steps__work_center')
            .order_by('-order_date', '-pk')
        )
        kept = [mo for mo in all_mos if mo.sales_order_id not in set(need_ids)]
        all_mos = kept + opened
        usable = [
            mo for mo in all_mos
            if mo.status not in (
                SxProductionOrder.STATUS_DRAFT,
                SxProductionOrder.STATUS_CANCELLED,
            )
        ]
    closes = team_job_closes(slug=(team.get('slug') or ''), mo_ids=[mo.pk for mo in usable])
    _fill_maps(all_mos)
    for row in report.rows:
        if row.is_total or not row.order_id:
            continue
        group = by_order.get(row.order_id) or []
        if not group:
            owned = any_by_order.get(row.order_id) or []
            if not owned:
                row.progress_block = 'chưa có lệnh sản xuất'
            elif all(mo.status == SxProductionOrder.STATUS_CANCELLED for mo in owned):
                row.progress_block = 'lệnh đã hủy'
            elif not any(
                mo.status not in (SxProductionOrder.STATUS_DRAFT, SxProductionOrder.STATUS_CANCELLED)
                for mo in owned
            ):
                row.progress_block = 'lệnh còn nháp'
            else:
                row.progress_block = 'không có công đoạn của tổ'
            continue
        open_mos = [
            mo for mo in group
            if mo.status != SxProductionOrder.STATUS_DONE and mo.pk not in closes
        ]
        chosen = open_mos or group[:1]
        row.progress_block = ''
        row.entry_mos = [
            {
                'id': mo.pk,
                'code': mo.code,
                'label': (mo.product_name or mo.product_code or mo.code),
            }
            for mo in chosen
        ]
        row.entry_mos_json = json.dumps(row.entry_mos, ensure_ascii=False)
        order = group[0].sales_order if group[0].sales_order_id else None
        if order is not None and order.plan_status == order.PLAN_DONE:
            continue
        row.closed_mos = [
            {'id': mo.pk, 'code': mo.code}
            for mo in group
            if mo.pk in closes
        ]
    return report


def expand_team_report_by_order(report: TeamOutputReport) -> TeamOutputReport:
    """Trang từng tổ: mỗi đơn một dòng. Modal ngày vẫn là danh sách đủ đơn của tổ."""
    for stage in report.stages:
        expanded: list[TeamOutputRow] = []
        for source in stage.rows:
            order_rows = _order_rows_from_team(source)
            if not order_rows:
                continue
            expanded.append(TeamOutputRow(
                slug=source.slug,
                label='Cộng',
                stage_slug=source.stage_slug,
                stage_label=source.stage_label,
                color_slug=source.color_slug,
                cells=source.cells,
                total=source.total,
                plan_total=source.plan_total,
                active_days=source.active_days,
                avg_per_day=source.avg_per_day,
                best_day=source.best_day,
                best_qty=source.best_qty,
                low_day=source.low_day,
                low_qty=source.low_qty,
                is_total=True,
            ))
            expanded.extend(order_rows)
        stage.rows = expanded
    report.by_order = True
    report.active_team_count = sum(1 for row in report.rows if row.total > 0 and not row.is_total)
    report.idle_teams = []
    return report


_WD_VN = ('Thứ 2', 'Thứ 3', 'Thứ 4', 'Thứ 5', 'Thứ 6', 'Thứ 7', 'Chủ nhật')


def _modal_rank_map(rows: list[TeamOutputRow]) -> dict[tuple[str, date, int], int]:
    """Thứ tự đã lưu của modal ngày — không đọc plan_rank lộ trình."""
    from san_xuat.hub_models import SxTeamDayModalRank

    slugs = {row.slug for row in rows}
    days = {
        cell.day
        for row in rows
        for cell in row.cells
        if cell.plan_orders or cell.actual_orders
    }
    if not slugs or not days:
        return {}
    out: dict[tuple[str, date, int], int] = {}
    for slug, day, oid, sort in SxTeamDayModalRank.objects.filter(
        team_slug__in=slugs,
        plan_date__in=days,
    ).values_list('team_slug', 'plan_date', 'sales_order_id', 'sort_order'):
        out[(slug, day, int(oid))] = int(sort or 0)
    return out


def _day_detail_payload(rows: list[TeamOutputRow]) -> dict:
    """JSON cho modal: bấm một ngày của một tổ → các đơn cần làm ngày đó."""
    from san_xuat.services.products import product_gallery_map
    from san_xuat.templatetags.sx_format import format_sx_num

    ranks = _modal_rank_map(rows)
    extra_codes = [
        item.get('product_code') or ''
        for row in rows
        for cell in row.cells
        for item in cell.actual_orders
        if item['code'] not in {plan['code'] for plan in cell.plan_orders}
    ]
    galleries = product_gallery_map(extra_codes) if extra_codes else {}
    payload: dict[str, dict] = {}
    for row in rows:
        for cell in row.cells:
            if not cell.plan_orders and not cell.actual_orders:
                continue
            done_left = {item['code']: item for item in cell.actual_orders}
            orders = []
            for item in cell.plan_orders:
                done = done_left.pop(item['code'], None)
                orders.append({
                    'id': int(item.get('id') or 0),
                    'code': item['code'],
                    'priority': item.get('priority') or '',
                    'priority_label': item.get('priority_label') or '',
                    'product': item.get('product') or '',
                    'image': item.get('image') or '',
                    'images': item.get('images') or [],
                    'qty': format_sx_num(item['qty']),
                    'done': format_sx_num(done['qty'] if done else 0),
                    'due': item['due'],
                    'url': item['url'],
                })
            for item in done_left.values():
                urls = galleries.get((item.get('product_code') or '').casefold()) or []
                orders.append({
                    'id': int(item.get('id') or 0),
                    'code': item['code'],
                    'priority': item.get('priority') or '',
                    'priority_label': item.get('priority_label') or '',
                    'product': item.get('product') or '',
                    'image': urls[0] if urls else '',
                    'images': urls[:6],
                    'qty': format_sx_num(0),
                    'done': format_sx_num(item['qty']),
                    'due': item['due'],
                    'url': item['url'],
                })

            def _order_key(item: dict) -> tuple:
                oid = int(item.get('id') or 0)
                rank = ranks.get((row.slug, cell.day, oid))
                return (rank if rank else 10**9, item.get('code') or '')

            orders.sort(key=_order_key)
            payload[f'{row.slug}|{cell.day.isoformat()}'] = {
                'team': row.label,
                'slug': row.slug,
                'date': cell.day.strftime('%d/%m/%Y'),
                'iso': cell.day.isoformat(),
                'weekday': _WD_VN[cell.day.weekday()],
                'plan': format_sx_num(cell.plan or 0),
                'done': format_sx_num(cell.qty),
                'orders': orders,
            }
    return payload


def _unique_order_ids(raw_ids) -> list[int]:
    out: list[int] = []
    seen: set[int] = set()
    for raw in raw_ids or []:
        try:
            oid = int(raw)
        except (TypeError, ValueError):
            continue
        if oid <= 0 or oid in seen:
            continue
        seen.add(oid)
        out.append(oid)
    return out


def _merge_visible_order(existing: list[int], sequence: list[int]) -> list[int]:
    """Đổi chỗ các đơn đang hiện, giữ nguyên đơn không có trong danh sách."""
    vis = set(sequence)
    pending = iter(sequence)
    merged: list[int] = []
    for oid in existing:
        if oid in vis:
            merged.append(next(pending, oid))
        else:
            merged.append(oid)
    seen = set(merged)
    merged.extend(oid for oid in sequence if oid not in seen)
    return merged


def reorder_team_day_orders(*, team_slug: str, plan_date: date, ordered_ids) -> int:
    """Ghi thứ tự modal một ngày của một tổ. Không đụng plan_rank của lộ trình."""
    from django.db import transaction

    from san_xuat.hub_models import SxSalesOrder, SxTeamDayModalRank
    from san_xuat.services.planning import PlanningError

    slug = (team_slug or '').strip().lower()
    if (
        not slug
        or len(slug) > 40
        or not slug.replace('-', '').replace('_', '').isalnum()
    ):
        raise PlanningError('Tổ không hợp lệ.')
    if not isinstance(plan_date, date):
        raise PlanningError('Ngày không hợp lệ.')
    sequence = _unique_order_ids(ordered_ids)
    if len(sequence) < 2:
        return 0
    known = set(
        SxSalesOrder.objects.filter(pk__in=sequence, is_demo=False).values_list('pk', flat=True)
    )
    sequence = [oid for oid in sequence if oid in known]
    if len(sequence) < 2:
        return 0

    with transaction.atomic():
        locked = list(
            SxTeamDayModalRank.objects.select_for_update().filter(
                team_slug=slug,
                plan_date=plan_date,
            ).order_by('sort_order', 'id')
        )
        by_order = {row.sales_order_id: row for row in locked}
        merged = _merge_visible_order([row.sales_order_id for row in locked], sequence)
        updated = 0
        for index, oid in enumerate(merged, start=1):
            row = by_order.get(oid)
            if row is None:
                SxTeamDayModalRank.objects.create(
                    team_slug=slug,
                    plan_date=plan_date,
                    sales_order_id=oid,
                    sort_order=index,
                )
                updated += 1
                continue
            if row.sort_order != index:
                row.sort_order = index
                row.save(update_fields=['sort_order'])
                updated += 1
        return updated


def teams_stage_order(teams) -> list[SimpleNamespace]:
    seen: list[str] = []
    for col in teams:
        s = (col.meta.get('stage_slug') or '').strip().lower()
        if s not in seen:
            seen.append(s)
    return [SimpleNamespace(slug=s) for s in seen]


def export_team_output_xlsx(report: TeamOutputReport):
    """Xuất bảng Tổ × ngày ra Excel (cùng bố cục tab Tổ SX)."""
    from io import BytesIO

    from django.http import HttpResponse
    from openpyxl import Workbook
    from openpyxl.cell.cell import MergedCell
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = 'San luong to'
    thin = Side(style='thin', color='CBD5E1')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    head_fill = PatternFill('solid', fgColor='F1F5F9')
    off_fill = PatternFill('solid', fgColor='F8FAFC')
    plan_fill = PatternFill('solid', fgColor='DBEAFE')
    state_font = {
        'met': Font(color='047857', bold=True),
        'under': Font(color='B45309', bold=True),
        'miss': Font(color='B91C1C', bold=True),
    }
    bold = Font(bold=True)
    center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    left = Alignment(horizontal='left', vertical='center')
    num_fmt = '#,##0'

    ws.cell(row=1, column=1, value=(
        f'Sản lượng tổ theo ngày {report.date_from:%d/%m/%Y} – {report.date_to:%d/%m/%Y}'
    )).font = Font(bold=True, size=14)

    summary_heads = ['Cộng'] if report.by_order else ['Cả kỳ', '% KH', 'TB/ngày', 'Cao nhất', 'Số ngày']
    n_days = len(report.days)
    name_col = 2 if report.by_order else 1
    kind_col = 4 if report.by_order else 2
    first_day_col = kind_col + 1
    head_row = 3
    if report.by_order:
        ws.cell(row=head_row, column=1, value='STT')
        ws.cell(row=head_row, column=3, value='Mức độ')
    ws.cell(row=head_row, column=name_col, value='Đơn' if report.by_order else 'Tổ')
    ws.cell(row=head_row, column=kind_col, value='')
    for i, c in enumerate(report.days):
        ws.cell(row=head_row, column=first_day_col + i, value=f'{c.day:%d/%m}')
    sum_col = first_day_col + n_days
    for j, label in enumerate(summary_heads):
        ws.cell(row=head_row, column=sum_col + j, value=label)
    last_col = sum_col + len(summary_heads) - 1
    for col in range(1, last_col + 1):
        cell = ws.cell(row=head_row, column=col)
        cell.font = bold
        cell.fill = head_fill
        cell.alignment = center
        cell.border = border
    for i, c in enumerate(report.days):
        if not c.is_workday:
            ws.cell(row=head_row, column=first_day_col + i).fill = off_fill

    def _num(cell, v) -> None:
        cell.value = float(v) if v else None
        cell.number_format = num_fmt
        cell.alignment = center
        cell.border = border

    def _write_pair(r: int, label: str, cells: list[TeamOutputCell], *, total, plan_total, rate, avg, best, days, source: TeamOutputRow, stt: int) -> int:
        """Hai dòng: Dự kiến (trên) · Thực hiện (dưới), như lưới Lộ trình."""
        if report.by_order:
            ws.cell(row=r, column=1, value='' if source.is_total else stt)
            ws.merge_cells(start_row=r, start_column=1, end_row=r + 1, end_column=1)
            ws.cell(row=r, column=1).alignment = center
            ws.cell(row=r, column=1).font = bold
            prio = '' if source.is_total else (source.priority_label or '')
            if source.progress_block:
                prio = f'{prio} ({source.progress_block})'.strip()
            ws.cell(row=r, column=3, value=prio)
            ws.merge_cells(start_row=r, start_column=3, end_row=r + 1, end_column=3)
            ws.cell(row=r, column=3).alignment = center
        ws.cell(row=r, column=name_col, value=label)
        ws.merge_cells(start_row=r, start_column=name_col, end_row=r + 1, end_column=name_col)
        ws.cell(row=r, column=name_col).alignment = left
        ws.cell(row=r, column=name_col).font = bold
        ws.cell(row=r, column=kind_col, value='Dự kiến')
        ws.cell(row=r + 1, column=kind_col, value='Thực hiện')
        for i, c in enumerate(cells):
            col = first_day_col + i
            plan_cell = ws.cell(row=r, column=col)
            _num(plan_cell, c.plan)
            if c.plan:
                plan_cell.fill = plan_fill
            done_cell = ws.cell(row=r + 1, column=col)
            _num(done_cell, c.qty)
            if c.state:
                done_cell.font = state_font[c.state]
            if not c.is_workday and not c.plan:
                plan_cell.fill = off_fill
            if not c.is_workday and not c.qty:
                done_cell.fill = off_fill
        _num(ws.cell(row=r, column=sum_col), plan_total)
        _num(ws.cell(row=r + 1, column=sum_col), total)
        extras = [] if report.by_order else [(f'{rate}%' if rate is not None else None), avg, best, days]
        for j, v in enumerate(extras, start=1):
            ws.merge_cells(start_row=r, start_column=sum_col + j, end_row=r + 1, end_column=sum_col + j)
            cell = ws.cell(row=r, column=sum_col + j)
            if isinstance(v, Decimal):
                _num(cell, v)
            else:
                cell.value = v or None
                cell.alignment = center
                cell.border = border
        for rr in (r, r + 1):
            for col in range(1, last_col + 1):
                cell = ws.cell(row=rr, column=col)
                if isinstance(cell, MergedCell):
                    continue
                cell.border = border
        return r + 2

    r = head_row + 1
    stt = 0
    for row in report.rows:
        label = row.label
        if report.by_order and row.note and not row.is_total:
            label = f'{row.label} — {row.note}'
        if report.by_order and not row.is_total:
            stt += 1
        r = _write_pair(
            r,
            label,
            row.cells,
            total=row.total,
            plan_total=row.plan_total,
            rate=row.plan_rate,
            avg=row.avg_per_day,
            best=row.best_qty,
            days=row.active_days,
            source=row,
            stt=stt,
        )

    if report.by_order:
        ws.column_dimensions['A'].width = 8
        ws.column_dimensions['B'].width = 28
        ws.column_dimensions['C'].width = 12
        ws.column_dimensions['D'].width = 12
    else:
        ws.column_dimensions['A'].width = 26
        ws.column_dimensions['B'].width = 10
    for col in range(first_day_col, sum_col):
        ws.column_dimensions[get_column_letter(col)].width = 7.5
    for col in range(sum_col, last_col + 1):
        ws.column_dimensions[get_column_letter(col)].width = 11
    ws.row_dimensions[head_row].height = 30
    ws.freeze_panes = ws.cell(row=head_row + 1, column=first_day_col)

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    response = HttpResponse(
        output.read(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = (
        f'attachment; filename=San_luong_to_{report.date_from:%Y%m%d}_{report.date_to:%Y%m%d}.xlsx'
    )
    return response
