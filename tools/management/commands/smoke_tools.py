"""Kiểm tra nhanh toàn bộ công cụ — chạy: python manage.py smoke_tools"""

import io
import os
import tempfile

from django.conf import settings
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management.base import BaseCommand
from django.test import Client, override_settings
from django.urls import reverse
from docx import Document
from openpyxl import Workbook
from PIL import Image
from reportlab.pdfgen import canvas

from tools.catalog import PORTAL_TOOLS, get_portal_tool_groups
from tools.services import (
    _libreoffice_binary,
    apply_image_watermark,
    compress_image,
    convert_image_format,
    convert_office_to_pdf,
    convert_pdf_to_docx,
    generate_qr_image,
    is_background_removal_ready,
    remove_image_background,
)


class Command(BaseCommand):
    help = 'Smoke test tất cả công cụ portal (service + trang GET).'

    def add_arguments(self, parser):
        parser.add_argument(
            '--skip-rmbg',
            action='store_true',
            help='Bỏ qua xóa nền (chậm / cần model).',
        )

    def handle(self, *args, **options):
        results = []
        results.append(self._check_catalog())
        results.extend(self._check_services(skip_rmbg=options['skip_rmbg']))
        results.extend(self._check_pages())

        failed = [line for ok, line in results if not ok]
        passed = len(results) - len(failed)
        self.stdout.write('')
        for ok, line in results:
            style = self.style.SUCCESS if ok else self.style.ERROR
            self._safe_write(style(f"{'OK' if ok else 'FAIL'}  {line}"))

        self.stdout.write('')
        if failed:
            self._safe_write(self.style.ERROR(f'Smoke tools: {passed}/{len(results)} pass'))
            raise SystemExit(1)
        self._safe_write(self.style.SUCCESS(f'Smoke tools: {passed}/{len(results)} pass'))

    def _safe_write(self, msg):
        try:
            self.stdout.write(msg)
        except UnicodeEncodeError:
            self.stdout.write(msg.encode('ascii', errors='replace').decode('ascii'))

    def _check_catalog(self):
        groups = get_portal_tool_groups()
        total = sum(len(g['tools']) for g in groups)
        ok = total == len(PORTAL_TOOLS) == 10 and len(groups) == 3
        return ok, f'catalog: {len(groups)} groups, {total} tools'

    def _check_services(self, *, skip_rmbg: bool):
        lines = []

        pdf_buf = io.BytesIO()
        c = canvas.Canvas(pdf_buf)
        c.drawString(72, 720, 'Smoke PDF')
        c.save()
        uploaded = SimpleUploadedFile('smoke.pdf', pdf_buf.getvalue(), content_type='application/pdf')
        try:
            data, name = convert_pdf_to_docx(uploaded)
            lines.append((data.startswith(b'PK'), 'service PDF to Word'))
        except Exception as exc:
            lines.append((False, f'service PDF to Word ({exc})'))

        doc_buf = io.BytesIO()
        doc = Document()
        doc.add_paragraph('Smoke Word')
        doc.save(doc_buf)
        doc_upload = SimpleUploadedFile(
            'smoke.docx',
            doc_buf.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
        )
        if _libreoffice_binary():
            try:
                pdf_data, pdf_name = convert_office_to_pdf(doc_upload)
                lines.append((pdf_data.startswith(b'%PDF'), 'service Word to PDF'))
            except Exception as exc:
                lines.append((False, f'service Word to PDF ({exc})'))
        else:
            lines.append((True, 'service Word to PDF (skip - no LibreOffice)'))

        xls_buf = io.BytesIO()
        wb = Workbook()
        wb.active['A1'] = 'Smoke'
        wb.save(xls_buf)
        xls_upload = SimpleUploadedFile(
            'smoke.xlsx',
            xls_buf.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        if _libreoffice_binary():
            try:
                pdf_data, _ = convert_office_to_pdf(xls_upload)
                lines.append((pdf_data.startswith(b'%PDF'), 'service Excel to PDF'))
            except Exception as exc:
                lines.append((False, f'service Excel to PDF ({exc})'))
        else:
            lines.append((True, 'service Excel to PDF (skip)'))

        img_buf = io.BytesIO()
        Image.new('RGB', (240, 180), color=(30, 144, 255)).save(img_buf, format='JPEG')
        img_upload = SimpleUploadedFile('smoke.jpg', img_buf.getvalue(), content_type='image/jpeg')

        try:
            data, name, ctype = compress_image(img_upload, quality=70, max_width=200)
            lines.append((len(data) > 0 and 'image' in ctype, 'service compress image'))
        except Exception as exc:
            lines.append((False, f'service compress image ({exc})'))

        img_upload.seek(0)
        try:
            data, name, ctype = convert_image_format(img_upload, 'png')
            lines.append((data.startswith(b'\x89PNG'), 'service convert image format'))
        except Exception as exc:
            lines.append((False, f'service convert image format ({exc})'))

        img_upload.seek(0)
        try:
            data, name, ctype = apply_image_watermark(img_upload, text='Smoke', opacity=30)
            lines.append((data.startswith(b'\x89PNG'), 'service watermark'))
        except Exception as exc:
            lines.append((False, f'service watermark ({exc})'))

        try:
            png = generate_qr_image('https://justplay.vn/smoke')
            lines.append((png.startswith(b'\x89PNG'), 'service QR'))
        except Exception as exc:
            lines.append((False, f'service QR ({exc})'))

        if skip_rmbg:
            lines.append((True, 'service remove bg (skipped)'))
        elif is_background_removal_ready():
            img_upload.seek(0)
            try:
                data, name = remove_image_background(img_upload)
                lines.append((data.startswith(b'\x89PNG'), 'service remove bg'))
            except Exception as exc:
                lines.append((False, f'service remove bg ({exc})'))
        else:
            lines.append((True, 'service remove bg (skip - model not warm)'))

        return lines

    def _check_pages(self):
        from django.contrib.auth.models import User

        user, _ = User.objects.get_or_create(username='_smoke_tools', defaults={'is_active': True})
        user.set_password('smoke-pass-internal')
        user.save()

        allowed = list(settings.ALLOWED_HOSTS)
        if 'testserver' not in allowed:
            allowed.append('testserver')

        lines = []
        with override_settings(ALLOWED_HOSTS=allowed):
            client = Client(HTTP_HOST='testserver')
            client.login(username='_smoke_tools', password='smoke-pass-internal')

            home = client.get(reverse('home_portal'))
            lines.append((
                home.status_code == 200 and b'jp-home-tools-group' in home.content,
                'home page tool groups',
            ))

            for tool in PORTAL_TOOLS:
                resp = client.get(reverse(tool['url_name']))
                lines.append((resp.status_code == 200, f'GET {tool["url_name"]}'))

        return lines
